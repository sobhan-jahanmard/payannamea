"""Read supporting content as evidence, with stable, addressable source locators."""
from __future__ import annotations

import hashlib
from pathlib import Path


def extract_source(path: Path, source_id: str, media_dir: Path) -> dict:
    segments, images = [], []

    def add(locator, text, pages=None):
        if text.strip():
            segments.append({"id": f"{source_id}:{locator}", "text": text.strip(), **({"pages": pages} if pages else {})})

    suffix = path.suffix.lower()
    if suffix == ".docx":
        from docx import Document
        document = Document(path)
        # Include tables in document order, so a locator retains its context.
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        for n, element in enumerate(document.element.body, 1):
            if element.tag.endswith("}p"):
                add(f"block-{n}", Paragraph(element, document).text)
            elif element.tag.endswith("}tbl"):
                table = Table(element, document)
                add(f"table-{n}", "\n".join(" | ".join(cell.text for cell in row.cells) for row in table.rows))
        for n, shape in enumerate(document.inline_shapes, 1):
            blips = shape._inline.xpath(".//a:blip")
            if not blips:
                continue
            from docx.oxml.ns import qn
            part = document.part.related_parts.get(blips[0].get(qn("r:embed")))
            if part and Path(part.partname).suffix.lower() in {".png", ".jpg", ".jpeg"}:
                media_dir.mkdir(parents=True, exist_ok=True)
                target = media_dir / f"{source_id}-figure-{n}{Path(part.partname).suffix}"
                target.write_bytes(part.blob)
                images.append({"id": f"{source_id}:figure-{n}", "path": str(target), "source_id": source_id})
                parent = shape._inline
                while parent.getparent() is not None and parent.getparent() is not document.element.body:
                    parent = parent.getparent()
                if parent.getparent() is document.element.body:
                    position = list(document.element.body).index(parent) + 1
                    images[-1]["source_locator"] = f"{source_id}:{'table' if parent.tag.endswith('}tbl') else 'block'}-{position}"
    elif suffix == ".pdf":
        from pypdf import PdfReader
        for n, page in enumerate(PdfReader(path).pages, 1):
            add(f"page-{n}", page.extract_text() or "", [n])
        import fitz
        with fitz.open(path) as document:
            for n, page in enumerate(document, 1):
                for index, entry in enumerate(page.get_images(full=True), 1):
                    bitmap = document.extract_image(entry[0])
                    if bitmap and bitmap["ext"].lower() in {"png", "jpg", "jpeg"}:
                        media_dir.mkdir(parents=True, exist_ok=True)
                        target = media_dir / f"{source_id}-page-{n}-image-{index}.{bitmap['ext']}"
                        target.write_bytes(bitmap["image"])
                        images.append({"id": f"{source_id}:page-{n}-image-{index}", "path": str(target), "source_id": source_id,
                                       "source_locator": f"{source_id}:page-{n}", "pages": [n]})
    elif suffix == ".pptx":
        from pptx import Presentation
        for n, slide in enumerate(Presentation(path).slides, 1):
            text = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    text.append(shape.text)
                elif shape.has_table:
                    text.extend(" | ".join(cell.text for cell in row.cells) for row in shape.table.rows)
                elif shape.has_chart:
                    chart = shape.chart
                    categories = [category.label for category in chart.plots[0].categories]
                    text.extend(f"{series.name}: " + "; ".join(f"{label}={value}" for label, value in zip(categories, series.values)) for series in chart.series)
                elif shape.shape_type == 13:  # Native embedded picture.
                    image = shape.image
                    if image.ext.lower() in {"png", "jpg", "jpeg"}:
                        media_dir.mkdir(parents=True, exist_ok=True)
                        target = media_dir / f"{source_id}-slide-{n}-image-{shape.shape_id}.{image.ext}"
                        target.write_bytes(image.blob)
                        images.append({"id": f"{source_id}:slide-{n}-image-{shape.shape_id}", "path": str(target), "source_id": source_id, "source_locator": f"{source_id}:slide-{n}", "pages": [n]})
            if slide.has_notes_slide:
                text.append(slide.notes_slide.notes_text_frame.text)
            add(f"slide-{n}", "\n".join(text), [n])
    elif suffix in {".md", ".txt"}:
        for n, line in enumerate(path.read_text(encoding="utf-8-sig").splitlines(), 1):
            add(f"line-{n}", line)
    else:
        raise ValueError(f"Unsupported information source: {path.name}")
    if not segments:
        raise ValueError(f"No readable text in {path.name}; supply a text PDF or DOCX (scanned PDFs need OCR first).")
    return {"id": source_id, "name": path.name, "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "segments": segments, "images": images}
