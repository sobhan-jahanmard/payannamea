#!/usr/bin/env python3
from __future__ import annotations

import argparse
import importlib
import json
import os
import re
import subprocess
import sys
import time
from datetime import datetime, timezone
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from uuid import uuid4

from common.order_context import load, new_context, save
from common.config import Config, load_config
from common.helpers import write_text, write_json
from common.api import claim_oldest, heartbeat, record_run
from common.profiles import profile_for
from common.execution_trace import begin_attempt, begin_session, finish_attempt
from common.codex_history import begin_call, finish_call, remember_thread

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")


STEPS = [item["id"] for item in json.loads((Path(__file__).with_name("stages.json")).read_text(encoding="utf-8"))]
FAILURE_STEP = "11_handle_failure"
TOTAL_STEPS = len(STEPS) + 1
WORKER_LABEL = "Order workflow"
MAX_RECOVERY_CYCLES = 12


def selected_workspace(root: Path, continuing: bool) -> Path:
    """Use the same active run that Electron displays when resuming work."""
    if continuing:
        pointer = root / "active.json"
        if pointer.is_file():
            try:
                relative = json.loads(pointer.read_text(encoding="utf-8"))["path"]
                selected = (root / relative).resolve()
                if selected.is_relative_to(root.resolve()) and (selected / "order_context.json").is_file():
                    return selected
            except (OSError, ValueError, KeyError, TypeError):
                pass
    return root / "in_progress"

# A failed quality gate usually means that an earlier artifact needs correction,
# not that the order itself has failed.  Restart from the earliest owner of that
# artifact and replay every dependent step.
RECOVERY_RESTARTS = {
    "06_build_thesis_plan": "05c_research_evidence",
    "06a_plan_visuals": "06a_plan_visuals",
    "06b_generate_visual_assets": "06b_generate_visual_assets",
    "07a_validate_visual_placement": "07_generate_content",
    "08_review_content": "07_generate_content",
    "08b_audit_source_compliance": "07_generate_content",
    "09_package_docx": "07_generate_content",
    "09b_polish_cover": "09_package_docx",
    "10b_finalize_persian_pagination": "09_package_docx",
    "10c_verify_persian_pagination": "09_package_docx",
    "10_validate_and_publish": "09_package_docx",
}


def recovery_start(step_name: str, error: Exception | None = None) -> int:
    if error and is_file_access_error(error):
        # A locked output does not make any upstream artifact invalid.
        return STEPS.index(step_name)
    if step_name == "08_review_content" and error and any(
        marker in str(error) for marker in ("missing deliverable:", "direct/proxy result needs", "source lacks URL")
    ):
        return STEPS.index("05c_research_evidence")
    if step_name == "10_validate_and_publish" and error and (
        "Visual policy gate failed" in str(error) or "AI visual quality audit failed" in str(error)
    ):
        return STEPS.index("06a_plan_visuals")
    if step_name == "10_validate_and_publish" and error and any(
        marker in str(error) for marker in ("page window", "page count outside order budget", "page count exceeds order budget", "Page-budget")
    ):
        return STEPS.index("07_generate_content")
    return STEPS.index(RECOVERY_RESTARTS.get(step_name, step_name))


def error_fingerprint(step_name: str, error: Exception) -> str:
    # Normalize standalone counts, but keep identifiers such as V01 distinct.
    # Otherwise defects in different visuals are mistaken for one repeated bug.
    message = re.sub(r"(?<![a-z0-9])\d+(?![a-z0-9])", "#", str(error).lower())
    message = re.sub(r"\s+", " ", message).strip()
    return f"{step_name}:{message[:500]}"


def is_deferred_service_error(error: Exception) -> bool:
    """Do not mark an order failed when the model provider is temporarily unavailable."""
    message = str(error).lower()
    return any(marker in message for marker in (
        "usage limit", "purchase more credits", "try again at", "rate limit", "temporarily unavailable",
        "out of credits", "refill", "token_revoked", "invalidated oauth token", "missing bearer", "unauthorized",
        "pdf font verification failed", "required font is not installed",
    ))


def is_file_access_error(error: Exception) -> bool:
    """Wait for a file/permission problem instead of regenerating valid work."""
    if isinstance(error, PermissionError) or (isinstance(error, OSError) and getattr(error, "winerror", None) in {32, 33}):
        return True
    message = str(error).lower()
    return any(marker in message for marker in (
        "permission denied", "access is denied", "winerror 5", "sharing violation", "being used by another process", "file is locked for editing",
    ))


def preserve_unstructured_failure(context: dict[str, Any]) -> None:
    """Keep failures written by older runners before resume clears their error list."""
    if context.get("status") not in {"failed", "paused"}:
        return
    error = str((context.get("errors") or [""])[-1] or context.get("pause_reason") or "")
    if not error:
        return
    history = context.setdefault("recovery_history", [])
    if history and history[-1].get("error") == error:
        return
    nodes = context.get("execution_trace", {}).get("nodes", [])
    step = context.get("current_step") or (nodes[-1].get("step") if nodes else None)
    if step not in STEPS:
        step = next((stage for stage in STEPS if stage not in context.get("completed_steps", [])), None)
    if not step:
        return
    failure = {"failed_step": step, "restart_step": step, "error": error}
    if is_file_access_error(RuntimeError(error)):
        failure["kind"] = "file_access"
    history.append(failure)
    context["active_failure"] = failure


def has_declared_chart_figure(markdown: str, manifest: Path | None = None) -> bool:
    """Use the manifest's explicit kind; captions need not repeat the word chart."""
    if manifest and manifest.is_file():
        try:
            figures = json.loads(manifest.read_text(encoding="utf-8")).get("figures", [])
            if any(str(item.get("kind") or "").lower() in {"chart", "graph"} for item in figures if isinstance(item, dict)):
                return True
        except (OSError, json.JSONDecodeError):
            pass
    return bool(re.search(r"!\[(?:[^\]]*?(?:نمودار|گراف|chart|graph)[^\]]*)\]\([^)]+\)", markdown, re.I))


@dataclass
class Services:
    config: Config
    args: argparse.Namespace
    workspace: Path
    token_usage: dict[str, int] = None  # type: ignore[assignment]
    profile: Any = None
    recovery_feedback: dict[str, Any] | None = None
    context: dict[str, Any] | None = None
    active_attempt_id: str | None = None
    context_write_depth: int = 0

    def __post_init__(self) -> None:
        self.token_usage = {"inputTokens": 0, "cachedInputTokens": 0, "outputTokens": 0, "reasoningTokens": 0, "totalTokens": 0}
        self._automation_word_pids_at_start = self.automation_word_pids()

    @staticmethod
    def automation_word_pids() -> set[int]:
        """Return only headless Word COM servers, never interactive Word."""
        if os.name != "nt":
            return set()
        command = (
            "Get-CimInstance Win32_Process -Filter \"Name='WINWORD.EXE'\" | "
            "Where-Object { $_.CommandLine -match '/Automation\\s+-Embedding' } | "
            "ForEach-Object { $_.ProcessId }"
        )
        result = subprocess.run(
            ["powershell.exe", "-NoProfile", "-Command", command],
            stdout=subprocess.PIPE, stderr=subprocess.DEVNULL, text=True, encoding="utf-8", errors="replace", timeout=15, check=False,
        )
        return {int(value) for value in re.findall(r"^\s*(\d+)\s*$", result.stdout, re.M)}

    def cleanup_run(self) -> dict[str, Any]:
        """Release worker DOCX locks and remove only COM servers started by this run."""
        report: dict[str, Any] = {"released_docx_targets": [], "terminated_automation_word_pids": [], "errors": []}
        if os.name != "nt":
            return report
        for path in sorted(self.workspace.rglob("*.docx")):
            try:
                self.release_docx_lock(path)
                report["released_docx_targets"].append(str(path.relative_to(self.workspace)))
            except Exception as exc:
                report["errors"].append(f"{path.name}: {exc}")
        # Word COM processes frequently become children of Explorer, not Python.
        # Capture the baseline at startup and terminate only new, hidden
        # `/Automation -Embedding` servers created by this worker run.
        for pid in sorted(self.automation_word_pids() - self._automation_word_pids_at_start):
            result = subprocess.run(
                ["taskkill", "/PID", str(pid), "/T", "/F"],
                stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, check=False,
            )
            if result.returncode == 0:
                report["terminated_automation_word_pids"].append(pid)
        write_json(self.workspace / "reports" / "stage_checks" / "run_cleanup.json", report)
        return report

    def run_codex(self, prompt: str, target: Path) -> None:
        if target.exists() and target.read_text(encoding="utf-8").strip():
            return
        if self.args.dry_run:
            write_text(target, "# خروجی آزمایشی\n\nاین خروجی فقط برای dry-run ساخته شده است.\n")
            return
        if self.recovery_feedback:
            feedback = self.recovery_feedback
            prompt += (
                "\n\n## بازخورد اعتبارسنجی اجرای قبلی\n"
                f"مرحلهٔ خطادار: {feedback.get('failed_step', '')}\n"
                f"خطای دقیق: {str(feedback.get('error', ''))[:2000]}\n"
                "قبل از تولید خروجی، artifactهای موجود و علت خطا را بررسی کن؛ "
                "با کمترین تغییر لازم همان علت را رفع کن و داده یا منبع جعل نکن.\n"
            )
        operator_note = str(getattr(self.args, "operator_note", "") or "").strip()
        if operator_note:
            failure_context = str(getattr(self.args, "failure_context", "") or "").strip()
            prompt += (
                "\n\n## دستور اپراتور برای این اجرای مجدد\n"
                "این یادداشت را با توجه به قرارداد سفارش، فایل‌های موجود و علت خطای قبلی تحلیل کن. "
                "آن را کورکورانه اجرا نکن؛ راه‌حلی انتخاب کن که خطای قبلی تکرار نشود و قواعد قطعی پروژه نقض نشوند.\n"
                f"یادداشت اپراتور:\n{operator_note}\n"
            )
            if failure_context:
                prompt += f"\nخطای ثبت‌شده در اجرای قبلی:\n{failure_context}\n"
        step = str(self.context.get("current_step") or "") if self.context else ""
        persist_context = self.context_write_depth == 0
        call = begin_call(self.context, self.workspace, self.active_attempt_id, step, target, prompt, persist=persist_context) if self.context and self.active_attempt_id and step else None
        resumed_thread = str(call.get("thread_id") or "") if call else ""
        def command_for(thread_id: str) -> list[str]:
            command = [
                self.config.codex_bin, "exec", *(["resume"] if thread_id else []), "--model", self.config.codex_model,
                "--json", "--output-last-message", str(target),
                "--skip-git-repo-check", "--ignore-user-config",
                *([] if getattr(self.profile, "ORDER_TYPE", None) == "ارائه و پاورپوینت" else ["--enable", "browser_use"]),
                "--dangerously-bypass-approvals-and-sandbox",
                *([] if thread_id else ["-C", str(self.workspace)]),
                *([thread_id] if thread_id else []), "-",
            ]
            return ["cmd.exe", "/d", "/s", "/c", subprocess.list2cmdline(command)] if os.name == "nt" else command
        codex_env = os.environ.copy()
        if self.config.codex_home is not None:
            codex_env["CODEX_HOME"] = str(self.config.codex_home)
        try:
            result = subprocess.run(command_for(resumed_thread), cwd=str(self.workspace), env=codex_env, input=prompt, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, encoding="utf-8", errors="replace", check=False)
            missing_session = bool(resumed_thread and result.returncode and re.search(
                r"(?:thread|session|rollout).{0,90}(?:not found|does not exist|missing|could not be found|failed to load)",
                result.stdout, re.I))
            if missing_session and call and self.context:
                history = [item for node in self.context["execution_trace"]["nodes"]
                           for item in node.get("codex_calls", []) if item.get("scope") == step and item["id"] != call["id"]]
                digest = "\n".join(f"- {item.get('target')}: {str(item.get('error') or item.get('response') or '')[:600]}"
                                   for item in history[-6:])
                prompt += ("\n\n## سابقهٔ بازیابی‌شدهٔ این مرحله\n"
                           "فایل جلسهٔ قبلی Codex در این دستگاه موجود نیست. خروجی‌ها و گزارش‌های workspace را بررسی کن.\n"
                           + digest)
                write_text(self.workspace / call["prompt_path"], prompt)
                call.update(continued=False, resumed_from=resumed_thread, thread_id=None)
                self.context["codex_threads"].pop(step, None)
                if persist_context:
                    save(self.workspace, self.context)
                first_output = result.stdout
                result = subprocess.run(command_for(""), cwd=str(self.workspace), env=codex_env, input=prompt, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, encoding="utf-8", errors="replace", check=False)
                write_text(self.workspace / call["events_path"], first_output + "\n" + result.stdout)
        except Exception as exc:
            if call and self.context:
                finish_call(self.context, self.workspace, call, status="failed", error=str(exc), persist=persist_context)
            raise
        if call and self.context and not missing_session:
            write_text(self.workspace / call["events_path"], result.stdout)
        # CLI JSON events can contain the full contents of every file read by
        # the agent. Keep the worker console useful without echoing megabytes
        # of source material or private order data into the Electron log.
        final_messages: list[str] = []
        service_errors: list[str] = []
        for line in result.stdout.splitlines():
            try:
                event = json.loads(line)
            except json.JSONDecodeError:
                continue
            if call and self.context and event.get("type") == "thread.started":
                thread_id = event.get("thread_id")
                if isinstance(thread_id, str) and thread_id:
                    remember_thread(self.context, self.workspace, call, thread_id, persist=persist_context)
            item = event.get("item") if isinstance(event.get("item"), dict) else {}
            if event.get("type") == "item.completed" and item.get("type") == "agent_message":
                final_messages.append(str(item.get("text") or ""))
            if event.get("type") in {"error", "turn.failed"}:
                detail = event.get("error") or event.get("message") or ""
                if isinstance(detail, dict):
                    detail = detail.get("message") or str(detail)
                service_errors.append(str(detail))
        if final_messages:
            print("[Codex] " + final_messages[-1][:1500], flush=True)
        if service_errors:
            print("[Codex error] " + service_errors[-1][:1500], flush=True)
        # A single `codex exec` can emit several cumulative usage events. Keep
        # only the final/highest value for this invocation, then add it to the
        # order-wide usage below. Taking max directly on self.token_usage would
        # lose every invocation except the most expensive one.
        invocation_usage = {"inputTokens": 0, "cachedInputTokens": 0, "outputTokens": 0, "reasoningTokens": 0, "totalTokens": 0}
        for line in result.stdout.splitlines():
            try: event = json.loads(line)
            except json.JSONDecodeError: continue
            usage = event.get("usage") or event.get("data", {}).get("usage") or {}
            for key, usage_field in (("input_tokens", "inputTokens"), ("cached_input_tokens", "cachedInputTokens"), ("output_tokens", "outputTokens"), ("total_tokens", "totalTokens")):
                if isinstance(usage.get(key), int): invocation_usage[usage_field] = max(invocation_usage[usage_field], usage[key])
            detail = usage.get("output_tokens_details") or {}
            input_detail = usage.get("input_tokens_details") or {}
            if isinstance(input_detail.get("cached_tokens"), int): invocation_usage["cachedInputTokens"] = max(invocation_usage["cachedInputTokens"], input_detail["cached_tokens"])
            if isinstance(detail.get("reasoning_tokens"), int): invocation_usage["reasoningTokens"] = max(invocation_usage["reasoningTokens"], detail["reasoning_tokens"])
            if isinstance(usage.get("reasoning_output_tokens"), int): invocation_usage["reasoningTokens"] = max(invocation_usage["reasoningTokens"], usage["reasoning_output_tokens"])
        if not invocation_usage["totalTokens"]:
            invocation_usage["totalTokens"] = invocation_usage["inputTokens"] + invocation_usage["outputTokens"]
        for usage_field, amount in invocation_usage.items():
            self.token_usage[usage_field] += amount
        if result.returncode:
            message = f"Codex exited with code {result.returncode}: {(service_errors[-1] if service_errors else result.stdout[-500:]).strip()[:1000]}"
            if call and self.context:
                finish_call(self.context, self.workspace, call, status="failed", response=final_messages[-1] if final_messages else "", error=message, persist=persist_context)
            raise RuntimeError(message)
        deadline = time.monotonic() + 1800
        while time.monotonic() < deadline:
            if target.exists() and target.read_text(encoding="utf-8").strip():
                if call and self.context:
                    finish_call(self.context, self.workspace, call, status="succeeded", response=final_messages[-1] if final_messages else "", persist=persist_context)
                return
            time.sleep(5)
        message = f"Codex did not create {target.relative_to(self.workspace)} within 30 minutes"
        if call and self.context:
            finish_call(self.context, self.workspace, call, status="failed", error=message, persist=persist_context)
        raise RuntimeError(message)

    def render_pdf_pages(self, pdf: Path, output_dir: Path) -> list[Path]:
        """Render final pages without relying on a machine-specific Poppler PATH."""
        try:
            import pymupdf
        except ImportError as exc:
            raise RuntimeError("PDF visual audit requires PyMuPDF; install worker_plus/requirements.txt") from exc
        output_dir.mkdir(parents=True, exist_ok=True)
        for stale in output_dir.glob("page-*.png"):
            stale.unlink()
        try:
            document = pymupdf.open(pdf)
            zoom = 150 / 72
            pages = []
            for number, page in enumerate(document, start=1):
                output = output_dir / f"page-{number:03}.png"
                page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), alpha=False).save(output)
                pages.append(output)
            document.close()
        except Exception as exc:
            raise RuntimeError(f"PDF visual rendering failed: {exc}") from exc
        if not pages:
            raise RuntimeError("PDF visual rendering produced no pages")
        return pages

    def release_docx_lock(self, path: Path) -> None:
        """Close only a stale Word instance for this worker artifact."""
        if os.name != "nt" or not path.exists():
            return
        script = Path(__file__).resolve().parent / "scripts" / "release_docx_lock.ps1"
        result = subprocess.run(
            ["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass", "-STA", "-File", str(script), "-Path", str(path)],
            stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, encoding="utf-8", errors="replace", timeout=60, check=False,
        )
        if result.returncode:
            raise RuntimeError("Could not release the previous Word document lock: " + result.stdout.strip()[-500:])

    def write_docx(self, source: Path, output: Path, title: str, rules: dict[str, Any], order: dict[str, Any]) -> None:
        try:
            from docx import Document
            from docx.enum.section import WD_SECTION
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.enum.style import WD_STYLE_TYPE
            from docx.shared import Cm, Pt, RGBColor
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
        except ImportError as exc:
            raise RuntimeError("python-docx is required. Install worker requirements first.") from exc

        font = rules["font"]
        page = rules["page"]
        paragraph_rules = rules["paragraph"]

        def bidi(paragraph: Any) -> None:
            props = paragraph._p.get_or_add_pPr()
            if props.find(qn("w:bidi")) is None:
                props.append(OxmlElement("w:bidi"))

        def rtl_table(table: Any) -> None:
            """Make column order and cell flow RTL, not only the text in each cell."""
            props = table._tbl.tblPr
            if props.find(qn("w:bidiVisual")) is None:
                props.append(OxmlElement("w:bidiVisual"))

        def set_complex_script_size(run: Any, size_pt: float) -> None:
            """Word uses szCs, rather than sz, for Persian/Arabic glyphs."""
            run.font.size = Pt(size_pt)
            props = run._element.get_or_add_rPr()
            size_cs = props.find(qn("w:szCs"))
            if size_cs is None:
                size_cs = OxmlElement("w:szCs")
                props.append(size_cs)
            size_cs.set(qn("w:val"), str(round(size_pt * 2)))

        def apply_persian_run(run: Any, size_pt: float | None = None) -> None:
            run.font.name = font["persian"]
            run.font.color.rgb = RGBColor(0, 0, 0)
            props = run._element.get_or_add_rPr()
            fonts = props.get_or_add_rFonts()
            for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
                fonts.set(qn(f"w:{attribute}"), font["persian"])
            if props.find(qn("w:rtl")) is None:
                props.append(OxmlElement("w:rtl"))
            language = props.find(qn("w:lang"))
            if language is None:
                language = OxmlElement("w:lang")
                props.append(language)
            language.set(qn("w:bidi"), "fa-IR")
            if size_pt is not None:
                set_complex_script_size(run, size_pt)

        def normalize_persian(value: str) -> str:
            value = value.replace("ي", "ی").replace("ك", "ک").replace("ـ", "")
            # Content generators usually emit ASCII digits even in Persian prose.
            # Convert them while building the DOCX, so section labels, numbered
            # citations, captions, tables, and bibliography markers consistently
            # use Persian glyphs.  Arabic-Indic digits are normalised too because
            # they are a distinct Unicode digit set from Persian digits. URLs and
            # DOI identifiers remain machine-readable, because digit conversion
            # would make a copied source link invalid.
            digit_map = str.maketrans("0123456789٠١٢٣٤٥٦٧٨٩", "۰۱۲۳۴۵۶۷۸۹۰۱۲۳۴۵۶۷۸۹")
            protected = re.split(r"((?:https?://|www\.)\S+|\b10\.\d{4,9}/\S+)", value)
            value = "".join(part if index % 2 else part.translate(digit_map) for index, part in enumerate(protected))
            value = re.sub(r"\s+([،؛؟٪])", r"\1", value)
            value = re.sub(r"([،؛؟])(?=\S)", r"\1 ", value)
            value = re.sub(r"\s+", " ", value).strip()
            return value

        def add_page_border(section: Any) -> None:
            border = rules["page"].get("border", {})
            if not border.get("enabled", False): return
            sect_pr = section._sectPr
            borders = OxmlElement("w:pgBorders")
            borders.set(qn("w:offsetFrom"), "page")
            for edge in ("top", "left", "bottom", "right"):
                tag = OxmlElement(f"w:{edge}")
                tag.set(qn("w:val"), "single")
                tag.set(qn("w:sz"), str(border.get("size", 6)))
                tag.set(qn("w:space"), str(border.get("space", 16)))
                tag.set(qn("w:color"), str(border.get("color", "808080")))
                borders.append(tag)
            sect_pr.append(borders)

        def style(name: str, size: float, bold: bool = False, align: Any = WD_ALIGN_PARAGRAPH.JUSTIFY) -> Any:
            current = document.styles[name] if name in document.styles else document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            current.font.name = font["persian"]
            current._element.rPr.rFonts.set(qn("w:eastAsia"), font["persian"])
            current._element.rPr.rFonts.set(qn("w:cs"), font["persian"])
            current.font.size = Pt(size)
            current.font.bold = bold
            current.font.color.rgb = RGBColor(0, 0, 0)
            current.paragraph_format.alignment = align
            current.paragraph_format.line_spacing = paragraph_rules["line_spacing"]
            current.paragraph_format.space_before = Pt(0)
            current.paragraph_format.space_after = Pt(paragraph_rules["space_after_pt"])
            props = current._element.get_or_add_pPr()
            # Word's built-in Title style can carry a theme-coloured bottom rule.
            # It has no place in a formal Persian thesis cover.
            for border in props.findall(qn("w:pBdr")):
                props.remove(border)
            # Word's stock Title style enables contextual spacing, which silently
            # discards paragraph spacing between consecutive cover lines.
            for contextual_spacing in props.findall(qn("w:contextualSpacing")):
                props.remove(contextual_spacing)
            if props.find(qn("w:bidi")) is None:
                props.append(OxmlElement("w:bidi"))
            style_props = current._element.get_or_add_rPr()
            size_cs = style_props.find(qn("w:szCs"))
            if size_cs is None:
                size_cs = OxmlElement("w:szCs")
                style_props.append(size_cs)
            size_cs.set(qn("w:val"), str(round(size * 2)))
            return current

        document = Document()
        # Word otherwise chooses the UI/document locale for inherited paragraph
        # direction.  Declare Persian as the bidi language at document level too.
        settings = document.settings._element
        language = settings.find(qn("w:themeFontLang"))
        if language is None:
            language = OxmlElement("w:themeFontLang")
            settings.append(language)
        language.set(qn("w:bidi"), "fa-IR")
        section = document.sections[0]
        section.top_margin, section.bottom_margin = Cm(page["top_cm"]), Cm(page["bottom_cm"])
        section.right_margin, section.left_margin = Cm(page["right_cm"]), Cm(page["left_cm"])
        section.header_distance, section.footer_distance = Cm(1.2), Cm(1.2)
        # Keep the cover visually clean: page numbering begins with the body.
        section.different_first_page_header_footer = True
        add_page_border(section)
        # Persian prose should be fully justified; headings, covers, and table
        # cells intentionally use their own alignment.
        normal = style("Normal", font["body_pt"], align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        normal.paragraph_format.first_line_indent = Cm(paragraph_rules["first_line_cm"])
        style("TOC Entry", font["body_pt"], align=WD_ALIGN_PARAGRAPH.RIGHT)
        style("Figure Caption", rules["captions"]["font_pt"], align=WD_ALIGN_PARAGRAPH.CENTER)
        style("Table Caption", rules["captions"]["font_pt"], align=WD_ALIGN_PARAGRAPH.CENTER)
        style("Title", font["heading_1_pt"] + 2, True, WD_ALIGN_PARAGRAPH.CENTER)
        style("Heading 1", font["heading_1_pt"], True, WD_ALIGN_PARAGRAPH.CENTER)
        style("Heading 2", font["heading_2_pt"], True, WD_ALIGN_PARAGRAPH.RIGHT)
        style("Heading 3", font["heading_3_pt"], True, WD_ALIGN_PARAGRAPH.RIGHT)
        for heading in ("Heading 1", "Heading 2", "Heading 3"):
            document.styles[heading].paragraph_format.space_before = Pt(18)
            document.styles[heading].paragraph_format.space_after = Pt(10)
            document.styles[heading].paragraph_format.keep_with_next = True

        # Modelled on the Ferdowsi University thesis template: institution and
        # faculty, degree/field, title, author, then supervisors.  Typography and
        # white space make those levels visibly distinct on the cover.
        cover = rules.get("cover", {})
        faculty = order.get("faculty") or ""
        department = order.get("department") or ""
        faculty_line = " — ".join(part for part in (f"دانشکده {faculty}" if faculty else "", f"گروه {department}" if department else "") if part)
        research_report = order.get("order_type") == "تحقیق دانشگاهی"
        degree_line = (f"گزارش تحقیق درس {order.get('course_name') or ''}" if research_report
                       else f"پروژه طراحی برای دریافت درجه {order.get('degree') or 'کارشناسی'}")
        if order.get("field_of_study") and not research_report:
            degree_line += f" در رشته {order['field_of_study']}"
        cover_lines = [
            (order.get("university") or "", cover.get("institution_pt", 21), True, cover.get("institution_after_pt", 32)),
            (faculty_line, cover.get("metadata_pt", 14), False, cover.get("faculty_after_pt", 52)),
            (degree_line, cover.get("degree_pt", 15), True, cover.get("degree_after_pt", 48)),
            (title, cover.get("title_pt", 28), True, cover.get("title_after_pt", 54)),
            (order.get("student_name") or "", cover.get("student_pt", 19), True, cover.get("student_after_pt", 48)),
            (f"استاد درس: {order['instructor_name']}" if research_report and order.get("instructor_name") else f"استاد راهنما: {order['advisor_name']}" if not research_report and order.get("advisor_name") else "", cover.get("supervisor_pt", 13), False, cover.get("supervisor_after_pt", 22)),
            (f"استاد مشاور: {order['consultant_name']}" if not research_report and order.get("consultant_name") else "", cover.get("supervisor_pt", 13), False, cover.get("supervisor_after_pt", 22)),
        ]
        for cover_index, (text, size, bold, space_after) in enumerate(cover_lines):
            if text:
                p = document.add_paragraph(style="Title")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.line_spacing = cover.get("line_spacing", 1.5)
                p.paragraph_format.space_before = Pt(
                    cover.get("top_space_pt", 72) if cover_index == 0
                    else cover.get("student_block_before_pt", 100) if text == (order.get("student_name") or "")
                    else 0
                )
                p.paragraph_format.space_after = Pt(space_after)
                run = p.add_run(normalize_persian(text))
                run.font.bold = bold
                apply_persian_run(run, size)
                bidi(p)
        document.add_page_break()

        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer_run = footer.add_run()
        field_begin, field_instr, field_separator, field_result, field_end = (OxmlElement("w:fldChar"), OxmlElement("w:instrText"), OxmlElement("w:fldChar"), OxmlElement("w:t"), OxmlElement("w:fldChar"))
        field_begin.set(qn("w:fldCharType"), "begin")
        field_instr.set(qn("xml:space"), "preserve")
        field_instr.text = "PAGE \\* MERGEFORMAT"
        field_separator.set(qn("w:fldCharType"), "separate")
        field_result.text = "۱"
        field_end.set(qn("w:fldCharType"), "end")
        footer_run._r.extend([field_begin, field_instr, field_separator, field_result, field_end])
        apply_persian_run(footer_run)
        bidi(footer)

        lines = source.read_text(encoding="utf-8").splitlines()
        # A content-repair agent may legitimately remove the Markdown TOC while
        # fixing another issue. The DOCX must still have a real, clickable TOC;
        # derive one from the document's actual headings instead of failing a
        # packaging-only quality gate three times.
        has_markdown_toc = any(normalize_persian(item.strip()) == "# فهرست مطالب" for item in lines)
        if not has_markdown_toc:
            lines = ["# فهرست مطالب", "", *lines]
        index = 0
        chapter_count = 0
        redundant_cover_headings = {"عنوان پژوهش", "عنوان پایان نامه", "عنوان پایان‌نامه", "عنوان پروژه"}
        heading_anchors: dict[int, str] = {}
        heading_entries: list[tuple[int, str, str]] = []
        figure_entries: list[tuple[str, str]] = []
        table_entries: list[tuple[str, str]] = []
        caption_anchors: dict[int, str] = {}
        anchor_number = 1

        def index_label(value: str) -> str:
            # Keep the full source and caveat on the body caption, but put only
            # its short identifying title in the list of figures/tables.
            title_only = re.split(r"\s*(?:[—–-]\s*منبع\s*:|؛|؛\s*منبع\s*:)", normalize_persian(value), maxsplit=1)[0]
            if len(title_only) <= 105:
                return title_only
            return title_only[:105].rsplit(" ", 1)[0].rstrip(" ،؛:") + "…"

        for source_index, source_line in enumerate(lines):
            candidate = normalize_persian(source_line.strip())
            prefix = next((value for value in ("### ", "## ", "# ") if candidate.startswith(value)), None)
            if prefix is not None:
                heading_text = normalize_persian(candidate[len(prefix):])
                if heading_text not in redundant_cover_headings and heading_text != "فهرست مطالب":
                    anchor = f"toc_{anchor_number}"
                    heading_anchors[source_index] = anchor
                    if prefix != "### ":
                        heading_entries.append((1 if prefix == "# " else 2, heading_text, anchor))
                    anchor_number += 1
            image = re.fullmatch(r"!\[(.+?)\]\(([^)]+)\)", source_line.strip())
            if image:
                anchor = f"figure_{anchor_number}"
                figure_entries.append((index_label(image.group(1)), anchor))
                caption_anchors[source_index] = anchor
                anchor_number += 1
            elif re.match(r"^جدول\s+[0-9۰-۹]+", candidate):
                anchor = f"table_{anchor_number}"
                table_entries.append((index_label(candidate), anchor))
                caption_anchors[source_index] = anchor
                anchor_number += 1

        bookmark_id = 1

        def add_bookmark(paragraph: Any, anchor: str) -> None:
            nonlocal bookmark_id
            start = OxmlElement("w:bookmarkStart")
            start.set(qn("w:id"), str(bookmark_id))
            start.set(qn("w:name"), anchor)
            end = OxmlElement("w:bookmarkEnd")
            end.set(qn("w:id"), start.get(qn("w:id")))
            paragraph._p.insert(0, start)
            paragraph._p.append(end)
            bookmark_id += 1

        def add_toc_link(paragraph: Any, text: str, anchor: str) -> None:
            run = paragraph.add_run(text)
            apply_persian_run(run)
            hyperlink = OxmlElement("w:hyperlink")
            hyperlink.set(qn("w:anchor"), anchor)
            hyperlink.set(qn("w:history"), "1")
            paragraph._p.remove(run._r)
            hyperlink.append(run._r)
            paragraph._p.append(hyperlink)
        def add_page_reference(paragraph: Any, anchor: str) -> None:
            page_run = paragraph.add_run()
            begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
            instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = f" PAGEREF {anchor} \\h "
            separate = OxmlElement("w:fldChar"); separate.set(qn("w:fldCharType"), "separate")
            result = OxmlElement("w:t"); result.text = "۱"
            end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
            page_run._r.extend([begin, instr, separate, result, end])
            apply_persian_run(page_run)

        def add_index_table(entries: list[tuple[str, str]], levels: list[int] | None = None) -> None:
            # A fixed right title column and left page-number column prevent
            # Word's bidirectional tab stops from putting numbers mid-line.
            table = document.add_table(rows=0, cols=2)
            table.autofit = False
            rtl_table(table)
            table.columns[0].width = Cm(13.0)
            table.columns[1].width = Cm(2.0)
            for row_index, (label, anchor) in enumerate(entries):
                title_cell, number_cell = table.add_row().cells
                title_cell.width, number_cell.width = Cm(13.0), Cm(2.0)
                title_paragraph = title_cell.paragraphs[0]
                title_paragraph.style = document.styles["TOC Entry"]
                # Word mirrors physical alignment in an RTL table: LEFT puts
                # the title against the right edge of this wide cell.
                title_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                title_paragraph.paragraph_format.first_line_indent = Cm(0)
                title_paragraph.paragraph_format.right_indent = Cm(0.35 if levels and levels[row_index] == 2 else 0)
                title_paragraph.paragraph_format.space_after = Pt(3)
                title_paragraph.paragraph_format.line_spacing = 1.15
                bidi(title_paragraph)
                add_toc_link(title_paragraph, label, anchor)
                page_paragraph = number_cell.paragraphs[0]
                page_paragraph.style = document.styles["TOC Entry"]
                page_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                page_paragraph.paragraph_format.first_line_indent = Cm(0)
                page_paragraph.paragraph_format.space_after = Pt(3)
                page_paragraph.paragraph_format.line_spacing = 1.15
                bidi(page_paragraph)
                add_page_reference(page_paragraph, anchor)

        while index < len(lines):
            raw_line = lines[index]
            line = normalize_persian(raw_line.strip())
            if not line:
                index += 1
                continue
            # The worker builds the cover itself.  A generated Markdown title
            # block would otherwise become an empty duplicate title page.
            if line.startswith("# ") and normalize_persian(line[2:]) in redundant_cover_headings:
                title_line_index = index + 1
                while title_line_index < len(lines) and not lines[title_line_index].strip():
                    title_line_index += 1
                # This is always a redundant Markdown cover block: the official
                # cover above already carries the order title. Do not compare
                # strings here, because model output can vary only in diacritics.
                index = title_line_index + 1 if title_line_index < len(lines) else index + 1
                continue
            if line == "# فهرست مطالب":
                toc_heading = document.add_paragraph("فهرست مطالب", style="Heading 1")
                toc_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
                # Paragraph direction alone is insufficient for the strict RTL
                # audit: the heading's own Persian run must carry w:rtl too.
                for toc_run in toc_heading.runs:
                    apply_persian_run(toc_run)
                if chapter_count and not research_report:
                    toc_heading.paragraph_format.page_break_before = True
                chapter_count += 1
                index += 1
                while index < len(lines):
                    toc_line = normalize_persian(lines[index].strip())
                    if toc_line.startswith("# "):
                        break
                    index += 1
                add_index_table([(label, anchor) for level, label, anchor in heading_entries], [level for level, _, _ in heading_entries])
                if figure_entries:
                    figures_heading = document.add_paragraph("فهرست شکل‌ها", style="Heading 1")
                    figures_heading.paragraph_format.page_break_before = not research_report
                    figures_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER; bidi(figures_heading)
                    for run in figures_heading.runs: apply_persian_run(run)
                    add_index_table(figure_entries)
                    chapter_count += 1
                if table_entries:
                    tables_heading = document.add_paragraph("فهرست جدول‌ها", style="Heading 1")
                    tables_heading.paragraph_format.page_break_before = not research_report
                    tables_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER; bidi(tables_heading)
                    for run in tables_heading.runs: apply_persian_run(run)
                    add_index_table(table_entries)
                    chapter_count += 1
                continue
            if line.startswith("|") and index + 1 < len(lines) and set(lines[index + 1].strip()) <= {"|", "-", ":", " "}:
                rows: list[list[str]] = []
                index += 2
                while index < len(lines) and lines[index].strip().startswith("|"):
                    rows.append([cell.strip() for cell in lines[index].strip().strip("|").split("|")])
                    index += 1
                headers = [normalize_persian(cell.strip()) for cell in line.strip("|").split("|")]
                table = document.add_table(rows=1, cols=len(headers))
                table.style = "Table Grid"
                rtl_table(table)
                for cell, value in zip(table.rows[0].cells, headers):
                    cell.text = value
                    for p in cell.paragraphs:
                        p.alignment = WD_ALIGN_PARAGRAPH.CENTER; bidi(p)
                        for r in p.runs: apply_persian_run(r, font["table_pt"]); r.font.bold = True
                for row in rows:
                    cells = table.add_row().cells
                    for cell, value in zip(cells, row):
                        cell.text = normalize_persian(value)
                        for p in cell.paragraphs:
                            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; bidi(p)
                            for r in p.runs: apply_persian_run(r, font["table_pt"])
                continue
            # `line` is normalized for Persian prose, but normalizing it also
            # transforms ASCII digits in asset filenames (V04.png -> V۰۴.png).
            # Parse Markdown image paths from the raw source instead.
            image_match = re.fullmatch(r"!\[(.+?)\]\(([^)]+)\)", raw_line.strip())
            if image_match:
                caption, relative_image = image_match.groups()
                image_path = (source.parent / relative_image).resolve()
                if not image_path.is_file():
                    raise RuntimeError(f"Markdown figure is missing: {relative_image}")
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = paragraph.add_run()
                run.add_picture(str(image_path), width=Cm(14.5))
                caption_paragraph = document.add_paragraph(normalize_persian(caption), style="Figure Caption")
                caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                caption_paragraph.paragraph_format.first_line_indent = Cm(0)
                bidi(caption_paragraph)
                for run in caption_paragraph.runs:
                    apply_persian_run(run, rules["captions"]["font_pt"])
                if index in caption_anchors:
                    add_bookmark(caption_paragraph, caption_anchors[index])
                index += 1
                continue
            if line.startswith("# "):
                paragraph = document.add_paragraph(line[2:], style="Heading 1")
                # Every top-level section starts on a fresh page.  The first one
                # already follows the cover's explicit page break.
                if chapter_count and not research_report:
                    paragraph.paragraph_format.page_break_before = True
                chapter_count += 1
                if index in heading_anchors:
                    add_bookmark(paragraph, heading_anchors[index])
            elif line.startswith("## "):
                paragraph = document.add_paragraph(line[3:], style="Heading 2")
                anchor = heading_anchors.get(index)
                if anchor:
                    add_bookmark(paragraph, anchor)
            elif line.startswith("### "):
                paragraph = document.add_paragraph(line[4:], style="Heading 3")
                anchor = heading_anchors.get(index)
                if anchor:
                    add_bookmark(paragraph, anchor)
            elif line.startswith(("- ", "* ")):
                paragraph = document.add_paragraph(line[2:], style="Normal")
                paragraph.paragraph_format.first_line_indent = Cm(0)
                paragraph.paragraph_format.right_indent = Cm(0.6)
            elif re.match(r"^جدول\s+[0-9۰-۹]+", line):
                paragraph = document.add_paragraph(line, style="Table Caption")
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.first_line_indent = Cm(0)
                if index in caption_anchors:
                    add_bookmark(paragraph, caption_anchors[index])
            else:
                paragraph = document.add_paragraph(line, style="Normal")
            if paragraph.style.name == "Normal":
                paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            elif paragraph.style.name == "Heading 1":
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            paragraph.paragraph_format.widow_control = True
            bidi(paragraph)
            for run in paragraph.runs:
                apply_persian_run(run)
            index += 1
        self.release_docx_lock(output)
        output.parent.mkdir(parents=True, exist_ok=True)
        document.save(output)

    def word_rtl_audit(self, docx: Path, audit_only: bool = True) -> dict[str, int]:
        if os.name != "nt":
            raise RuntimeError("actual Word RTL audit requires the Windows Worker runtime")
        self.release_docx_lock(docx)
        script = Path(__file__).resolve().parent / "scripts" / "enforce_word_rtl.ps1"
        command = [
            "powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass", "-STA", "-File", str(script), "-Path", str(docx)
        ]
        if audit_only:
            command.append("-AuditOnly")
        result = subprocess.run(command, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, encoding="utf-8", errors="replace", timeout=180, check=False)
        if result.returncode:
            raise RuntimeError("Word RTL audit failed: " + result.stdout.strip()[-500:])
        counts = re.findall(r"^\s*(\d+)\s*,\s*(\d+)\s*,\s*(\d+)\s*$", result.stdout, re.M)
        if not counts:
            raise RuntimeError("Word RTL audit returned no numeric result")
        right, center, invalid = (int(value) for value in counts[-1])
        return {"right_rtl": right, "center_rtl": center, "invalid_or_left": invalid}

    def enforce_word_rtl(self, docx: Path) -> None:
        audit = self.word_rtl_audit(docx, audit_only=False)
        if audit["invalid_or_left"]:
            raise RuntimeError(f"Word RTL enforcement left {audit['invalid_or_left']} invalid Persian paragraphs")

    def polish_cover(self, docx: Path, rules: dict[str, Any], order: dict[str, Any]) -> None:
        """Dedicated cover-layout pass, separate from content packaging."""
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        cover = rules.get("cover", {})
        values = {
            "institution": (cover.get("institution_pt", 21), True, cover.get("institution_after_pt", 32)),
            "faculty": (cover.get("metadata_pt", 14), False, cover.get("faculty_after_pt", 52)),
            "degree": (cover.get("degree_pt", 15), True, cover.get("degree_after_pt", 48)),
            "title": (cover.get("title_pt", 28), True, cover.get("title_after_pt", 54)),
            "student": (cover.get("student_pt", 19), True, cover.get("student_after_pt", 48)),
            "supervisor": (cover.get("supervisor_pt", 13), False, cover.get("supervisor_after_pt", 22)),
        }
        document = Document(docx)
        title_props = document.styles["Title"]._element.get_or_add_pPr()
        for contextual_spacing in title_props.findall(qn("w:contextualSpacing")):
            title_props.remove(contextual_spacing)
        for index, paragraph in enumerate(p for p in document.paragraphs if p.style.name == "Title" and p.text.strip()):
            text = paragraph.text.strip()
            role = "supervisor"
            if text == (order.get("university") or ""):
                role = "institution"
            elif text == (order.get("title") or ""):
                role = "title"
            elif text == (order.get("student_name") or ""):
                role = "student"
            elif text.startswith("دانشکده "):
                role = "faculty"
            elif text.startswith("پروژه طراحی برای دریافت درجه") or text.startswith("گزارش تحقیق درس "):
                role = "degree"
            size, bold, gap = values[role]
            paragraph.paragraph_format.line_spacing = cover.get("line_spacing", 1.5)
            paragraph.paragraph_format.space_before = Pt(
                cover.get("top_space_pt", 72) if index == 0
                else cover.get("student_block_before_pt", 100) if role == "student"
                else 0
            )
            paragraph.paragraph_format.space_after = Pt(gap)
            for run in paragraph.runs:
                run.font.size = Pt(size)
                run.font.bold = bold
                run.font.color.rgb = RGBColor(0, 0, 0)
                props = run._element.get_or_add_rPr()
                props.get_or_add_rFonts().set(qn("w:cs"), rules["font"]["persian"])
                size_cs = props.find(qn("w:szCs"))
                if size_cs is None:
                    size_cs = OxmlElement("w:szCs")
                    props.append(size_cs)
                size_cs.set(qn("w:val"), str(round(size * 2)))
        document.save(docx)

    def count_docx_pages(self, docx: Path) -> int:
        if os.name != "nt":
            raise RuntimeError("actual DOCX page counting requires the Windows Worker runtime")
        self.release_docx_lock(docx)
        script = Path(__file__).resolve().parent / "scripts" / "measure_pages.ps1"
        result = subprocess.run([
            "powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass", "-STA", "-File", str(script), "-Path", str(docx)
        ], stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, encoding="utf-8", errors="replace", timeout=120, check=False)
        if result.returncode:
            raise RuntimeError("Word page count failed: " + result.stdout.strip()[-500:])
        counts = re.findall(r"^\s*(\d+)\s*$", result.stdout, re.M)
        if not counts:
            raise RuntimeError("Word page count returned no numeric result")
        return int(counts[-1])

    def export_pdf(self, docx: Path, pdf: Path, font_name: str) -> None:
        if os.name != "nt":
            raise RuntimeError("PDF export requires the Windows Word runtime")
        script = Path(__file__).resolve().parent / "scripts" / "export_pdf.ps1"
        last_output = ""
        for attempt in range(1, 3):
            self.release_docx_lock(docx)
            self.release_docx_lock(pdf)
            if pdf.exists():
                pdf.unlink()
            process = subprocess.Popen(
                ["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass", "-STA", "-File", str(script), "-Path", str(docx), "-OutputPath", str(pdf), "-FontName", font_name],
                stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, encoding="utf-8", errors="replace",
            )
            try:
                output, _ = process.communicate(timeout=180)
            except subprocess.TimeoutExpired:
                subprocess.run(["taskkill", "/PID", str(process.pid), "/T", "/F"], stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, check=False)
                output, _ = process.communicate()
                last_output = output or "Word PDF export timed out"
                self.release_docx_lock(docx)
                if attempt == 1:
                    continue
                raise RuntimeError("Word PDF export timed out after retry")
            last_output = output
            if process.returncode == 0 and pdf.exists() and pdf.stat().st_size:
                self.verify_pdf_font(pdf, font_name)
                return
            self.release_docx_lock(docx)
        raise RuntimeError("Word PDF export failed: " + last_output[-500:])

    def verify_pdf_font(self, pdf: Path, font_name: str) -> None:
        """Require the requested thesis font in the actual exported PDF."""
        try:
            import pymupdf
        except ImportError as exc:
            raise RuntimeError("PDF font verification requires PyMuPDF; install worker_plus/requirements.txt") from exc

        def normalized(value: str) -> str:
            # PDF subset names are commonly written as ABCDEF+Font-Name.
            value = value.split("+", 1)[-1]
            return "".join(char.lower() for char in value if char.isalnum())

        expected = normalized(font_name)
        document = pymupdf.open(pdf)
        try:
            exported_fonts = {
                normalized(str(font[3]))
                for page in document
                for font in page.get_fonts(full=True)
            }
        finally:
            document.close()
        if not any(expected in exported or exported in expected for exported in exported_fonts):
            raise RuntimeError(
                f"PDF font verification failed: expected '{font_name}', found {sorted(exported_fonts)}"
            )

    def create_first_half_sample(self, docx: Path, sample_docx: Path) -> tuple[int, int]:
        """Save the actual first half of rendered Word pages as a separate DOCX."""
        self.release_docx_lock(docx)
        self.release_docx_lock(sample_docx)
        script = Path(__file__).resolve().parent / "scripts" / "create_first_half_sample.ps1"
        result = subprocess.run(
            ["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass", "-STA", "-File", str(script), "-Path", str(docx), "-OutputPath", str(sample_docx)],
            stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, encoding="utf-8", errors="replace", timeout=300, check=False,
        )
        counts = re.findall(r"^\s*(\d+)\s*,\s*(\d+)\s*$", result.stdout, re.M)
        if result.returncode or not sample_docx.exists() or not sample_docx.stat().st_size or not counts:
            raise RuntimeError("Word first-half sample creation failed: " + result.stdout[-500:])
        return tuple(map(int, counts[-1]))

    def validate_docx(self, source: Path, docx: Path, rules: dict[str, Any], report: Path, mode: str, order: dict[str, Any]) -> None:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        text = source.read_text(encoding="utf-8")
        # Word's effective paragraph formatting is authoritative. Repair it before
        # inspecting OOXML so a retry can converge instead of repeatedly rejecting
        # style-inherited RTL/alignment that Word renders correctly.
        word_rtl_audit = self.word_rtl_audit(docx)
        if word_rtl_audit["invalid_or_left"]:
            self.enforce_word_rtl(docx)
            word_rtl_audit = self.word_rtl_audit(docx)
        document = Document(docx)
        errors: list[str] = []
        warnings: list[str] = []
        # `--sample` controls submission status only; every run generates the
        # same complete package and its first-half excerpt.
        from common.page_budget import planned_page_budget
        budget = planned_page_budget(order, self.workspace / "planning" / "visual_plan.json", rules)
        min_words = min(rules["quality"]["min_full_words"], budget["minimum_content_words"]) if budget else rules["quality"]["min_full_words"]
        if len(re.findall(r"[آ-یA-Za-z]{2,}", text)) < min_words: errors.append(f"content below minimum word target ({min_words})")
        if not re.search(r"^# .*?(منابع|References)", text, re.M | re.I): errors.append("missing references section")
        citations = set(re.findall(r"\[([0-9۰-۹]+)\]", text))
        reference_numbers = set(re.findall(r"(?m)^\s*\[([0-9۰-۹]+)\]", text))
        if rules["quality"]["require_citations"] and not citations: errors.append("missing numbered in-text citations")
        missing_references = citations - reference_numbers
        if missing_references: errors.append("in-text citations missing from reference list: " + ", ".join(sorted(missing_references)))
        if any(marker in text.upper() for marker in ("TODO", "TBD", "[NEEDS", "لورم")): errors.append("placeholder marker found")
        source_table_count = sum(
            1 for index, line in enumerate(text.splitlines()[:-1])
            if line.strip().startswith("|") and bool(text.splitlines()[index + 1].strip()) and set(text.splitlines()[index + 1].strip()) <= {"|", "-", ":", " "}
        )
        if len(document.tables) < source_table_count: errors.append("not all Markdown tables were converted to DOCX tables")
        chart_required = bool(order.get("requires_charts"))
        manifest = self.workspace / "final" / "figures" / "manifest.json"
        if chart_required and not has_declared_chart_figure(text, manifest):
            errors.append("order requires a chart or graph but no declared chart figure exists")
        if document.styles["Normal"].font.name != rules["font"]["persian"]: errors.append("default Persian font was not applied")
        if len(document.paragraphs) < 12: errors.append("DOCX contains too few paragraphs")
        toc_present = any(paragraph.text.strip() == "فهرست مطالب" for paragraph in document.paragraphs)
        bookmarks = {item.get(qn("w:name")) for item in document._element.iter(qn("w:bookmarkStart"))}
        toc_links = [
            item.get(qn("w:anchor"))
            for item in document._element.iter(qn("w:hyperlink"))
            if item.get(qn("w:anchor"))
        ]
        if not toc_present:
            errors.append("missing table of contents")
        elif not toc_links:
            errors.append("table of contents has no clickable entries")
        elif any(anchor not in bookmarks for anchor in toc_links):
            errors.append("table of contents contains a broken internal link")
        required_pages = int(order.get("quantity_value") or 0) if order.get("quantity_type") == "pages" else 0
        actual_pages = self.count_docx_pages(docx)
        if required_pages and actual_pages < required_pages:
            errors.append(f"page count below order requirement ({actual_pages}/{required_pages})")
        if budget and actual_pages > budget["word_maximum"]:
            errors.append(f"page count exceeds order budget ({actual_pages}/{budget['word_maximum']} Word pages)")
        persian_runs = 0
        wrong_font_runs = 0
        missing_rtl_runs = 0
        bad_direction_paragraphs = 0
        non_justified_body_paragraphs = 0

        def style_property(paragraph: Any, tag: str) -> Any:
            style = paragraph.style
            while style is not None:
                props = style._element.pPr
                if props is not None:
                    value = props.find(qn(tag))
                    if value is not None:
                        return value
                style = style.base_style
            return None

        def style_complex_font(paragraph: Any) -> str | None:
            style = paragraph.style
            while style is not None:
                props = style._element.rPr
                fonts = props.rFonts if props is not None else None
                value = fonts.get(qn("w:cs")) if fonts is not None else None
                if value:
                    return value
                style = style.base_style
            return None
        # Centred text is only legitimate for the cover/chapter title and table
        # headers.  A centred or left-aligned body paragraph is a Word-layout bug.
        paragraphs_to_audit = [
            (paragraph, paragraph.style.name in ("Title", "Heading 1", "Figure Caption"), paragraph.style.name == "Normal")
            for paragraph in document.paragraphs
        ]
        paragraphs_to_audit.extend(
            (paragraph, row_index == 0, False)
            for table in document.tables
            for row_index, row in enumerate(table.rows)
            for cell in row.cells
            for paragraph in cell.paragraphs
        )
        for paragraph, center_allowed, justify_required in paragraphs_to_audit:
            contains_persian = any(re.search(r"[آ-ی]", run.text) for run in paragraph.runs)
            if not contains_persian:
                continue
            props = paragraph._p.pPr
            if (props is None or props.find(qn("w:bidi")) is None) and style_property(paragraph, "w:bidi") is None:
                bad_direction_paragraphs += 1
            effective_alignment = paragraph.alignment if paragraph.alignment is not None else paragraph.style.paragraph_format.alignment
            allowed_alignments = (WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.JUSTIFY)
            if center_allowed:
                allowed_alignments += (WD_ALIGN_PARAGRAPH.CENTER,)
            if effective_alignment not in allowed_alignments:
                bad_direction_paragraphs += 1
            if justify_required and effective_alignment != WD_ALIGN_PARAGRAPH.JUSTIFY:
                non_justified_body_paragraphs += 1
            for run in paragraph.runs:
                if not re.search(r"[آ-ی]", run.text):
                    continue
                persian_runs += 1
                run_props = run._element.rPr
                run_fonts = run_props.rFonts if run_props is not None else None
                effective_font = run_fonts.get(qn("w:cs")) if run_fonts is not None else style_complex_font(paragraph)
                if effective_font != rules["font"]["persian"]:
                    wrong_font_runs += 1
                if run_props is None or run_props.find(qn("w:rtl")) is None:
                    missing_rtl_runs += 1
        if wrong_font_runs: errors.append(f"{wrong_font_runs} Persian runs do not use the required font")
        if missing_rtl_runs: errors.append(f"{missing_rtl_runs} Persian runs are missing RTL direction")
        if non_justified_body_paragraphs:
            errors.append(f"{non_justified_body_paragraphs} Persian body paragraphs are not justified")
        if word_rtl_audit["invalid_or_left"]:
            errors.append(f"Word reports {word_rtl_audit['invalid_or_left']} Persian paragraphs as LTR or left-aligned")
        elif bad_direction_paragraphs:
            warnings.append(
                f"{bad_direction_paragraphs} Persian paragraphs rely on Word/style formatting; "
                "Word's effective RTL audit passed"
            )
        cover_paragraphs = [paragraph for paragraph in document.paragraphs if paragraph.style.name == "Title" and paragraph.text.strip()]
        cover_spacing_issues = 0
        if len(cover_paragraphs) < 2:
            cover_spacing_issues += 1
        else:
            for cover_index, paragraph in enumerate(cover_paragraphs):
                fmt = paragraph.paragraph_format
                style_fmt = paragraph.style.paragraph_format
                effective_line_spacing = fmt.line_spacing if fmt.line_spacing is not None else style_fmt.line_spacing
                # Word removes redundant line-spacing XML from single-line cover
                # paragraphs when saving. Missing XML is therefore not a visual
                # spacing failure; an explicit value must still be in range.
                line_spacing_ok = effective_line_spacing is None or 1.4 <= float(effective_line_spacing) <= 1.7
                after = fmt.space_after if fmt.space_after is not None else style_fmt.space_after
                before = fmt.space_before if fmt.space_before is not None else style_fmt.space_before
                after_pt = after.pt if after is not None else 0
                before_pt = before.pt if before is not None else 0
                if not line_spacing_ok or after_pt < 8 or (cover_index == 0 and before_pt < 60):
                    cover_spacing_issues += 1
        section_headings = [paragraph for paragraph in document.paragraphs if paragraph.style.name == "Heading 1"]
        missing_section_breaks = sum(
            1 for paragraph in section_headings[1:]
            if not paragraph.paragraph_format.page_break_before
        ) if order.get("order_type") != "تحقیق دانشگاهی" else 0
        if cover_spacing_issues:
            errors.append(f"cover page spacing audit failed ({cover_spacing_issues} issue(s))")
        cover_sizes = [run.font.size.pt for paragraph in cover_paragraphs for run in paragraph.runs if run.font.size is not None]
        title_cover = next((paragraph for paragraph in cover_paragraphs if paragraph.text.strip() == (order.get("title") or "")), None)
        student_cover = next((paragraph for paragraph in cover_paragraphs if paragraph.text.strip() == (order.get("student_name") or "")), None)
        title_size = max((run.font.size.pt for run in title_cover.runs if run.font.size is not None), default=0) if title_cover else 0
        if not cover_sizes or max(cover_sizes) - min(cover_sizes) < 8 or title_size != max(cover_sizes):
            errors.append("cover typography audit failed: title hierarchy or font-size contrast is insufficient")
        if student_cover is None or (student_cover.paragraph_format.space_before.pt if student_cover.paragraph_format.space_before else 0) < 60:
            errors.append("cover layout audit failed: student and supervisor block is not positioned in the lower cover area")
        if missing_section_breaks:
            errors.append(f"top-level sections missing required page break ({missing_section_breaks})")
        if rules.get("requires_manual_guideline_review"): warnings.append("uploaded guideline text was unreadable; fallback rules used pending manual/OCR review")
        result = {"passed": not errors, "errors": errors, "warnings": warnings, "word_target": min_words, "required_pages": required_pages or None, "actual_pages": actual_pages, "paragraphs": len(document.paragraphs), "citations": len(citations), "source_tables": source_table_count, "docx_tables": len(document.tables), "font": rules["font"]["persian"], "toc_audit": {"present": toc_present, "clickable_entries": len(toc_links), "broken_entries": sum(anchor not in bookmarks for anchor in toc_links)}, "rtl_audit": {"persian_runs": persian_runs, "wrong_font_runs": wrong_font_runs, "missing_rtl_runs": missing_rtl_runs, "bad_direction_paragraphs": bad_direction_paragraphs, "non_justified_body_paragraphs": non_justified_body_paragraphs, "word": word_rtl_audit}, "layout_audit": {"cover_paragraphs": len(cover_paragraphs), "cover_spacing_issues": cover_spacing_issues, "cover_font_sizes_pt": cover_sizes, "cover_title_size_pt": title_size, "top_level_sections": len(section_headings), "missing_section_breaks": missing_section_breaks}}
        write_json(report, result)
        if errors: raise RuntimeError("DOCX quality gate failed: " + "; ".join(errors))


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Worker Plus shared order workflow")
    parser.add_argument("--sample", action="store_true", help="Set the final order status to customer-sample approval; output generation remains complete")
    parser.add_argument("--order-id", "--order_id", dest="order_id", help="Force-claim this order and start from a fresh workspace")
    parser.add_argument("--redo", action="store_true", help="Allow reclaiming a specific order (implicit with --order-id)")
    parser.add_argument("--resume", action="store_true", help="Resume workspace/in_progress from its saved context")
    parser.add_argument("--dry-run", action="store_true", help="Avoid Codex and create a tiny test source")
    parser.add_argument("--offline", action="store_true", help="Run all steps with a local mock order; no backend changes")
    parser.add_argument("--publish-local", metavar="ORDER_ID", help="Upload the completed local package for this order without rebuilding it or calling AI")
    parser.add_argument("--repackage", action="store_true", help="Restart at Step 12 (Package DOCX) using the saved source, without invoking Codex")
    parser.add_argument("--step", type=int, choices=range(1, len(STEPS) + 1), help="Resume from a terminal step number; the final step validates, visually audits, and publishes the package")
    parser.add_argument("--auto-repair", action="store_true", help="Diagnose a saved failure, test an isolated stage-code repair if needed, then rerun the selected step")
    note_source = parser.add_mutually_exclusive_group()
    note_source.add_argument("--operator-note", default="", help="Operator guidance that every AI-assisted step must analyze and apply during this run")
    note_source.add_argument("--operator-note-stdin", action="store_true", help="Read operator guidance from standard input, as sent by the desktop app")
    return parser


def read_operator_note(args: argparse.Namespace, stream: Any) -> str:
    if args.operator_note_stdin:
        # Electron writes UTF-8 bytes. On Windows, a redirected Python stdin
        # can use the local ANSI code page and decode continuation bytes as
        # surrogate escapes, which cannot be saved to UTF-8 JSON later.
        raw = stream.read(262_145)
        if len(raw) > 262_144:
            raise SystemExit("Operator note exceeds 262,144 bytes")
        try:
            note = raw.decode("utf-8")
        except UnicodeDecodeError as exc:
            raise SystemExit("Operator note must be valid UTF-8") from exc
        if len(note) > 65_536:
            raise SystemExit("Operator note exceeds 65,536 characters")
        return note.strip()
    return str(args.operator_note or "").strip()


def header(index: int, title: str, context: dict[str, Any]) -> None:
    clean_title = re.sub(r"\s+", " ", title)
    print("-" * 64, flush=True)
    print(f"Worker Plus | {context.get('order', {}).get('order_type') or WORKER_LABEL}", flush=True)
    print(f"Order: {context.get('order_id') or 'pending'} | Mode: {context['mode']}", flush=True)
    print(f"Step {index:02}/{TOTAL_STEPS:02}: {clean_title}", flush=True)
    print("-" * 64, flush=True)


def main() -> None:
    args = build_parser().parse_args()
    args.operator_note = read_operator_note(args, sys.stdin.buffer)
    if args.auto_repair and not args.step:
        raise SystemExit("--auto-repair requires --step for a saved failed order")
    config = load_config()
    workspace = selected_workspace(config.workspace_root, bool(args.resume or args.step or args.repackage or args.publish_local))
    if args.publish_local:
        if args.resume or args.step or args.repackage or args.order_id or args.offline or args.sample or args.dry_run or args.redo or args.operator_note:
            raise SystemExit("--publish-local cannot be combined with run options.")
        from common.local_publish import publish_local
        publish_local(config, workspace, args.publish_local)
        return
    workspace.mkdir(parents=True, exist_ok=True)
    if args.order_id or (args.step == 1):
        # A fresh run uses the legacy preparation path, while the previously
        # selected isolated run remains on disk and can be inspected later.
        write_json(config.workspace_root / "active.json", {"path": str(workspace.relative_to(config.workspace_root)).replace("\\", "/")})
    if args.repackage and (args.resume or args.order_id or args.step):
        raise SystemExit("--repackage uses the saved workspace; do not combine it with --resume or --order-id.")
    if args.step and args.order_id:
        raise SystemExit("--step resumes the saved workspace; do not combine it with --order-id.")
    if args.resume and args.order_id:
        raise SystemExit("--resume cannot be combined with --order-id; a specific order always starts fresh.")
    args.redo = args.redo or bool(args.order_id) or args.repackage
    if args.repackage:
        context = load(workspace)
        preserve_unstructured_failure(context)
        if not context.get("order_id"):
            raise SystemExit("No saved order is available to repackage.")
        if not args.offline:
            claimed = claim_oldest(config, context["order_id"], redo=True)
            context["order"] = claimed["customerInput"]
        context["status"] = "in_progress"
        context["errors"] = []
        # The package step is the first step that can safely rebuild Word/PDF
        # outputs while retaining the reviewed source and all intake artifacts.
        context["completed_steps"] = STEPS[:STEPS.index("09_package_docx")]
    elif args.step:
        if args.step == 1:
            context = new_context(workspace, "sample" if args.sample else "full")
        else:
            context = load(workspace)
            preserve_unstructured_failure(context)
            if not context.get("order_id"):
                raise SystemExit("No saved order is available; start with a normal run before using --step.")
            # A resumed sample publish must honour an explicit --sample flag.
            # Previously the saved full-mode context silently won, so
            # `--step 16 --sample` published as a full run despite the CLI flag.
            if args.sample:
                context["mode"] = "sample"
            # Re-claiming and heartbeating keep the backend state correct even when
            # the expensive intake/source stages are intentionally skipped.
            if not args.offline:
                claimed = claim_oldest(config, context["order_id"], redo=True)
                context["order"] = claimed["customerInput"]
                heartbeat(config, context["order_id"], f"Worker Plus resumed from step {args.step}.")
            context["status"] = "in_progress"
            context["errors"] = []
            context.pop("active_failure", None)
            context.pop("pause_reason", None)
            context["completed_steps"] = STEPS[:args.step - 1]

            # A deliberate re-run from generation must not silently reuse a stale
            # model output.  Starting at packaging/validation leaves source data
            # untouched for visual-only corrections.
            if args.step <= STEPS.index("07_generate_content") + 1:
                for stale in (workspace / "final" / "deliverable_source.md", workspace / "final" / "sample_source.md"):
                    if stale.exists():
                        stale.unlink()
                for stale in (workspace / "drafts").glob("continuation_*.md"):
                    stale.unlink()
            if args.step <= STEPS.index("09_package_docx") + 1:
                for stale in (workspace / "final" / "deliverable.docx", workspace / "final" / "sample.docx"):
                    if stale.exists():
                        stale.unlink()
    else:
        context = load(workspace) if args.resume else new_context(workspace, "sample" if args.sample else "full")
        preserve_unstructured_failure(context)
    if args.resume:
        context["status"] = "in_progress"
        context["errors"] = []
        context.pop("active_failure", None)
        context.pop("pause_reason", None)
    args.operator_note = args.operator_note.strip()
    latest_recovery = (context.get("recovery_history") or [{}])[-1]
    args.failure_context = str(latest_recovery.get("error") or ((context.get("errors") or [""])[-1])).strip()
    context["operator_note"] = args.operator_note
    if args.operator_note:
        context.setdefault("operator_note_history", []).append({
            "created_at": datetime.now(timezone.utc).isoformat(),
            "from_step": STEPS[args.step - 1] if args.step else ("resume" if args.resume else STEPS[0]),
            "note": args.operator_note,
            "failure_context": args.failure_context or None,
        })
    save(workspace, context)
    services = Services(config, args, workspace)
    services.context = context
    if args.resume or args.step:
        next_step = STEPS[args.step - 1] if args.step else next(
            (step for step in STEPS if step not in context["completed_steps"]), None
        )
        if latest_recovery.get("error") and latest_recovery.get("restart_step") == next_step:
            services.recovery_feedback = latest_recovery
    trace_session = begin_session(context, STEPS, "manual" if args.step else "repackage" if args.repackage else "resume" if args.resume else "run")
    save(workspace, context)
    if args.auto_repair and services.recovery_feedback and not is_file_access_error(RuntimeError(args.failure_context)):
        from common.auto_repair import repair_worker_code
        trace = context["execution_trace"]
        failed_attempt = next((node for node in reversed(trace["nodes"]) if node["status"] in {"failed", "paused"} and node["step"] == latest_recovery.get("failed_step")), None)
        repair_event = {"id": str(uuid4()), "step": next_step, "after_attempt_id": failed_attempt["id"] if failed_attempt else None,
                        "status": "running", "thread_id": None, "note": args.operator_note,
                        "started_at": datetime.now(timezone.utc).isoformat(), "finished_at": None, "files": [], "reason": "", "turns": []}
        trace.setdefault("repair_events", []).append(repair_event)
        save(workspace, context)

        def remember_repair_thread(thread_id: str) -> None:
            threads = context.setdefault("code_repair_threads", {})
            if thread_id:
                threads[next_step] = thread_id
            else:
                threads.pop(next_step, None)
            repair_event["thread_id"] = thread_id or None
            save(workspace, context)
        def record_repair_turn(number: int, prompt: str, output: str, exit_code: int) -> None:
            prefix = f"reports/codex_calls/repair-{repair_event['id']}-{number}"
            write_text(workspace / f"{prefix}.prompt.md", prompt)
            write_text(workspace / f"{prefix}.events.jsonl", output or "")
            repair_event["turns"].append({"prompt_path": f"{prefix}.prompt.md", "events_path": f"{prefix}.events.jsonl",
                                          "exit_code": exit_code})
            save(workspace, context)
        try:
            decision = repair_worker_code(Path(__file__).resolve().parents[1], workspace,
                                          str(context.get("order", {}).get("order_type") or ""),
                                          next_step, {**latest_recovery, "operator_note": args.operator_note}, config.codex_bin,
                                          config.codex_model, config.codex_home,
                                          thread_id=context.get("code_repair_threads", {}).get(next_step),
                                          on_thread=remember_repair_thread, on_turn=record_repair_turn)
        except Exception as error:
            decision = {"status": "skipped", "reason": str(error)[:1000]}
        repair_event.update(status=decision["status"], thread_id=context.get("code_repair_threads", {}).get(next_step),
                            files=decision.get("files", []), reason=decision.get("reason", ""),
                            finished_at=datetime.now(timezone.utc).isoformat())
        save(workspace, context)
        decision.update({"order_id": context.get("order_id"), "stage": next_step,
                         "thread_id": repair_event["thread_id"],
                         "checked_at": datetime.now(timezone.utc).isoformat()})
        write_json(workspace / "reports" / "stage_checks" / "auto_code_repair.json", decision)
        print(f"Auto code repair: {decision['status']} — {decision.get('reason', ', '.join(decision.get('files', [])))[:500]}", flush=True)
        if decision["status"] not in {"applied", "no_code_change"}:
            context.update(status="paused", current_step=None, errors=[], active_failure=latest_recovery,
                           pause_reason="بررسی و اصلاح کامل نشد؛ جزئیات مانع را در گزارش اصلاح کد ببینید.")
            save(workspace, context)
            print("Result: PAUSED — code repair was blocked or unverified; the failed stage was not replayed.", flush=True)
            return
    if context.get("order"):
        services.profile = profile_for(context["order"])

    # Keep the note preflight inside the step's normal progress/error handling.
    # Previously it ran before current_step was set, making a live run look idle
    # and allowing preflight exceptions to bypass recovery reporting entirely.
    preflight_step = next((step for step in STEPS if step not in context["completed_steps"]), None)
    pending_note_preflight = bool(
        args.operator_note and context.get("completed_steps") and preflight_step
        and not is_file_access_error(RuntimeError(args.failure_context))
    )

    terminal_error: Exception | None = None
    recovery_history: list[dict[str, Any]] = context.setdefault("recovery_history", [])
    repeated_errors: dict[str, int] = {}
    # Keep the active failure streak across manual retries from the desktop UI.
    previous_fingerprint = None
    for failure in reversed(recovery_history):
        latest_attempt = next((node for node in reversed(context["execution_trace"]["nodes"])
                               if node["step"] == failure.get("failed_step")), None)
        if latest_attempt and latest_attempt["status"] == "succeeded":
            break
        fingerprint = error_fingerprint(str(failure.get("failed_step") or ""), RuntimeError(str(failure.get("error") or "")))
        if previous_fingerprint and fingerprint != previous_fingerprint:
            break
        repeated_errors[fingerprint] = repeated_errors.get(fingerprint, 0) + 1
        previous_fingerprint = fingerprint
    try:
        cursor = 0
        recovery_cycles = 0
        while cursor < len(STEPS):
            name = STEPS[cursor]
            if name in context["completed_steps"]:
                cursor += 1
                continue
            module = importlib.import_module(f"common.steps.{name}")
            presentation_order = context.get("order", {}).get("order_type") == "ارائه و پاورپوینت"
            if presentation_order:
                from presentation.workflow import LABELS as presentation_labels, run_step as run_presentation_step
            header(cursor + 1, presentation_labels.get(name, module.TITLE) if presentation_order else module.TITLE, context)
            context["current_step"] = name
            trace_attempt = begin_attempt(context, STEPS, name, trace_session)
            services.active_attempt_id = trace_attempt
            try:
                save(workspace, context)
                if pending_note_preflight and name == preflight_step:
                    analysis_report = workspace / "reports" / "operator_note_analysis.md"
                    analysis_report.parent.mkdir(parents=True, exist_ok=True)
                    if analysis_report.exists():
                        analysis_report.unlink()
                    preflight_prompt = f"""پیش از اجرای مجدد مرحله `{name}`، نقش تحلیل‌گر بازیابی workflow را انجام بده.

فایل `order_context.json`، گزارش‌های `reports/` و artifactهای مرتبط موجود در workspace را بخوان. علت خطای قبلی و یادداشت اپراتور که در انتهای این prompt آمده را تحلیل کن. اگر برای جلوگیری از تکرار خطا لازم است یکی از artifactهای موجود اصلاح شود، همان فایل را مستقیم و با کمترین تغییر لازم ویرایش کن. validatorها را حذف یا دور نزن، نتیجه جعلی PASS نساز، داده یا منبع جعل نکن و فایل‌های سالم نامرتبط را تغییر نده.

در خروجی Markdown کوتاه بنویس: علت محتمل، فایل‌های بررسی‌شده، تغییرهای انجام‌شده و اینکه مرحله `{name}` اکنون چه چیزی را باید اعتبارسنجی کند."""
                    print(f"Preparing step {cursor + 1:02}: analyzing operator note and prior failure.", flush=True)
                    if presentation_order:
                        from presentation.recovery import guard_inputs
                        preflight_prompt = f"""پیش از ادامه مرحله `{name}`، فقط `order_context.json` و گزارش خطای قبلی را بخوان و علت و پیشنهاد اصلاح را گزارش کن. هیچ فایل، منبع، locator، متن نقل‌شده یا پیش‌نویسی را ویرایش نکن. اصلاح اسلاید در مرحله تولید و پس از اعتبارسنجی انجام می‌شود. پاسخ Markdown کوتاه باشد."""
                        with guard_inputs(services):
                            services.run_codex(preflight_prompt, analysis_report)
                    else:
                        services.run_codex(preflight_prompt, analysis_report)
                    context["artifacts"]["operator_note_analysis"] = str(analysis_report.relative_to(workspace))
                    save(workspace, context)
                    pending_note_preflight = False
                if not presentation_order or not run_presentation_step(name, context, services):
                    module.run(context, services)
            except Exception as exc:
                deferred = is_deferred_service_error(exc) or is_file_access_error(exc)
                finish_attempt(context, trace_attempt, "paused" if deferred else "failed", str(exc))
                recovery_cycles += 1
                fingerprint = error_fingerprint(name, exc)
                repeated_errors[fingerprint] = repeated_errors.get(fingerprint, 0) + 1
                restart = recovery_start(name, exc)
                if presentation_order and not is_file_access_error(exc):
                    # Content/layout faults return to the slide model; renderer
                    # outages stay at rendering and never rebuild thesis prose.
                    presentation_restart = {
                        "07a_validate_visual_placement": "07_generate_content",
                        "08_review_content": "07_generate_content",
                        "08b_audit_source_compliance": "07_generate_content",
                        "09_package_docx": "07_generate_content",
                        "09b_polish_cover": "09_package_docx",
                        "09c_package_interim_reports": "09c_package_interim_reports",
                        "10b_finalize_persian_pagination": "10b_finalize_persian_pagination",
                        "10c_verify_persian_pagination": "09_package_docx",
                        "10_validate_and_publish": "07_generate_content",
                    }
                    restart = STEPS.index(presentation_restart.get(name, name))
                    if getattr(exc, "restart_step", None) in STEPS:
                        restart = STEPS.index(exc.restart_step)
                    if getattr(exc, "report_path", None) and not getattr(exc, "repair_slides", None):
                        # An audit without a located defect must clarify its
                        # report before any healthy slide is regenerated.
                        restart = STEPS.index(name)
                event = {
                    "cycle": recovery_cycles,
                    "failed_step": name,
                    "restart_step": STEPS[restart],
                    "same_error_count": repeated_errors[fingerprint],
                    "error": str(exc),
                }
                if presentation_order:
                    for field in ("report_path", "issues", "repair_slides", "repair_kind"):
                        if getattr(exc, field, None):
                            event[field] = getattr(exc, field)
                    context["active_failure"] = event
                if is_file_access_error(exc):
                    event["kind"] = "file_access"
                recovery_history.append(event)
                services.recovery_feedback = event
                context.setdefault("errors", []).append(str(exc))
                context["status"] = "in_progress"
                context["current_step"] = None
                if is_deferred_service_error(exc) or is_file_access_error(exc):
                    # Service availability and file access require an external
                    # change. Preserve validated work and pause on this step.
                    context["deferred_error"] = event
                    context["active_failure"] = event
                    context["status"] = "paused"
                    context["pause_reason"] = str(exc)
                    context["errors"] = []
                    save(workspace, context)
                    if is_file_access_error(exc):
                        print("Result: PAUSED — output file is locked or access is denied; close the file or resolve permissions, then resume this same step. Validated upstream work is preserved.", flush=True)
                    else:
                        print("Result: PAUSED — required service or runtime is unavailable; order remains in_progress.", flush=True)
                    return
                # Invalidate the failed artifact and everything derived from it.
                if presentation_order and getattr(exc, "validated_checkpoint", False):
                    context["completed_steps"] = STEPS[:restart]
                context["completed_steps"] = [step for step in context["completed_steps"] if STEPS.index(step) < restart]
                save(workspace, context)
                # Three identical failures are useful evidence, but still leave
                # room for broader replay cycles before declaring a terminal fault.
                must_stop = getattr(exc, "retryable", True) is False or (presentation_order and name in {"03_check_intake", "05a_analyze_customer_sources"}) or "made no progress" in str(exc).lower() or recovery_cycles >= MAX_RECOVERY_CYCLES or repeated_errors[fingerprint] >= (3 if presentation_order else 6)
                print(
                    f"Result: {'STOP' if must_stop else 'RETRY'} — {exc}\nRecovery {recovery_cycles}/{MAX_RECOVERY_CYCLES}: "
                    + (f"after fixing the reported issue, resume manually from {restart + 1:02} ({STEPS[restart]})" if must_stop
                       else f"restarting from {restart + 1:02} ({STEPS[restart]})"),
                    flush=True,
                )
                if must_stop:
                    terminal_error = exc
                    break
                cursor = restart
                continue
            if name not in context["completed_steps"]:
                context["completed_steps"].append(name)
            finish_attempt(context, trace_attempt, "succeeded")
            repeated_errors = {key: count for key, count in repeated_errors.items() if not key.startswith(name + ":")}
            context.pop("active_failure", None)
            context.pop("deferred_error", None)
            context.pop("pause_reason", None)
            services.recovery_feedback = None
            context["current_step"] = None
            context["status"] = "in_progress" if cursor + 1 < len(STEPS) else context["status"]
            save(workspace, context)
            print("Result: PASS", flush=True)
            cursor += 1
        if terminal_error is not None:
            raise terminal_error
        # Recovered errors are diagnostic history, not a failed run.
        context["errors"] = []
        if context.get("status") in {"dry_run_complete", "worker_done_pending_approval", "sample_pending_customer_approval", "completed"}:
            context.pop("operator_note", None)
            context.pop("active_failure", None)
            context.pop("pause_reason", None)
        save(workspace, context)
    except Exception as exc:
        context.setdefault("errors", []).append(str(exc))
        save(workspace, context)
        module = importlib.import_module(f"common.steps.{FAILURE_STEP}")
        header(len(STEPS) + 1, module.TITLE, context)
        try:
            module.run(context, services)
            save(workspace, context)
        finally:
            print(f"Result: FAIL — {exc}", flush=True)
        raise SystemExit(1) from exc
    finally:
        try:
            cleanup = services.cleanup_run()
            context["cleanup"] = cleanup
            save(workspace, context)
            print(
                "Cleanup: released "
                f"{len(cleanup['released_docx_targets'])} DOCX target(s); terminated "
                f"{len(cleanup['terminated_automation_word_pids'])} worker Word automation process(es).",
                flush=True,
            )
        except Exception as cleanup_error:
            print(f"Could not complete worker cleanup: {cleanup_error}", file=sys.stderr)
        if context.get("order_id") and not args.offline and not context.get("deferred_error"):
            try:
                errors = context.get("errors") or []
                record_run(config, context["order_id"], {"workerId": config.worker_id, "model": config.codex_model, "mode": context["mode"], "status": "completed" if not errors else "failed", "notes": errors[-1] if errors else None, **services.token_usage})
            except Exception as record_error:
                print(f"Could not record worker run usage: {record_error}", file=sys.stderr)


if __name__ == "__main__":
    main()
