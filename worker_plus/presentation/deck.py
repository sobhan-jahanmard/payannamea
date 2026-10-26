"""Portable, editable PPTX compiler. Content stays in a validated JSON slide model."""
from __future__ import annotations

import json
import math
import re
import unicodedata
from pathlib import Path


def normalized(text):
    return " ".join(str(text).translate(str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")).split())


def quoted_numbers(text: str) -> set[float]:
    """Read source numbers without losing a Unicode minus or merging a list.

    Word/PDF sources commonly use U+2212 MINUS SIGN rather than ASCII '-'.
    Decimal and thousands separators may also be Persian. Keep the sign in
    the numeric token so a positive chart value cannot pass for a loss.
    """
    text = normalized(text).translate(str.maketrans({
        "−": "-", "﹣": "-", "－": "-", "﹢": "+", "＋": "+", "٫": ".",
    }))
    tokens = re.finditer(r"(?P<prefix>[-+]?)\s*(?P<number>(?:\d{1,3}(?:[,٬]\d{3})+|\d+)(?:\.\d+)?)(?P<suffix>[-+](?!\s*\d))?", text)
    values = set()
    for token in tokens:
        value = float(re.sub(r"[,٬]", "", token["number"]))
        # RTL tables can place the minus after the digits, e.g. '۵٫۷۰−'.
        # A subtraction operator followed by another number is not a suffix.
        if token["prefix"] == "-" or token["suffix"] == "-":
            value = -value
        values.add(value)
    return values


def source_numbers(text: str) -> set[float]:
    """Include numbers written with a word for a minus sign in slide prose."""
    text = re.sub(r"(?i)(?<!\w)(?:منفی|minus|negative)\s+(?=[\d۰-۹٠-٩])", "-", str(text))
    return quoted_numbers(text)


def validate_spec(spec: dict, contract: dict, sources: dict, workspace: Path) -> None:
    slides = spec.get("slides", [])
    if len(slides) != contract["slide_count"]:
        raise ValueError(f"Slide count must be exactly {contract['slide_count']}; got {len(slides)}")
    locators = {s["id"]: s["text"] for source in sources["sources"] for s in source["segments"]}
    images = {image["id"]: image for source in sources["sources"] for image in source["images"]}
    analysis_packet = None
    if contract.get("analytical_assurance_version") == 1:
        from presentation.analysis import validate_packet
        from presentation.assurance import source_topics
        path = workspace / "planning/presentation_analysis.json"
        if not path.is_file():
            raise ValueError("Analytical packet missing; build evidence before slides")
        analysis_packet = validate_packet(json.loads(path.read_text(encoding="utf-8")), sources, source_topics(sources))
    if slides[0].get("section") != "title" or slides[0].get("layout") != "title":
        raise ValueError("First slide must be the title slide")
    for n, slide in enumerate(slides, 1):
        if analysis_packet is not None:
            from presentation.analysis import calculated_numbers
            analysis_numbers = calculated_numbers(analysis_packet, slide.get("analysis_ids", []))
        else:
            analysis_numbers = set()
        if n > 1 and (slide.get("section") == "title" or slide.get("layout") == "title"):
            raise ValueError(f"Slide {n}: only the first slide may use the cover layout")
        if slide.get("layout") not in {"title", "statement", "two_column", "image", "chart", "table", "comparison", "key_metric", "process", "timeline", "decision", "closing"}:
            raise ValueError(f"Slide {n}: unsupported layout")
        title = slide.get("title")
        points = slide.get("points", [])
        if not isinstance(title, str) or not title.strip() or len(title) > 140:
            raise ValueError(f"Slide {n}: title missing or too long")
        if not isinstance(points, list) or any(not isinstance(p, str) for p in points) or len(points) > contract["max_points"]:
            raise ValueError(f"Slide {n}: at most {contract['max_points']} concise points")
        if len(" ".join(points).split()) > contract["max_body_words"] or any(len(p) > 240 for p in points):
            raise ValueError(f"Slide {n}: overcrowded text; shorten or redistribute it")
        if not isinstance(slide.get("notes"), str) or not slide["notes"].strip():
            raise ValueError(f"Slide {n}: speaker notes required")
        if slide.get("section") not in contract["sections"] + ["agenda"]:
            raise ValueError(f"Slide {n}: unsupported narrative section")
        if slide.get("description") and (not isinstance(slide["description"], str) or len(slide["description"]) > 110):
            raise ValueError(f"Slide {n}: description must be a short, separate subtitle")
        evidence = slide.get("evidence", [])
        from presentation.provenance import validate_metadata_evidence
        validate_metadata_evidence(slide, sources, n)
        if slide.get("image_id") and slide["image_id"] not in images:
            raise ValueError(f"Slide {n}: unknown source image")
        if slide.get("section") not in {"title", "questions", "references", "agenda"} and not evidence:
            raise ValueError(f"Slide {n}: source evidence required")
        for reference in evidence:
            locator, quote = reference.get("locator"), normalized(reference.get("quote", ""))
            if locator not in locators or len(quote) < 8 or quote not in normalized(locators[locator]):
                raise ValueError(f"Slide {n}: source quote/locator mismatch: {locator}")
        if slide["layout"] == "image":
            image = images.get(slide.get("image_id"))
            if not image:
                raise ValueError(f"Slide {n}: image must come from the source inventory")
            path = (workspace / image["path"]).resolve()
            if not path.is_relative_to(workspace.resolve()) or not path.is_file():
                raise ValueError(f"Slide {n}: source image path is invalid")
            if not slide.get("caption") or not evidence:
                raise ValueError(f"Slide {n}: source image needs a caption and contextual evidence")
        if slide["layout"] == "chart":
            chart = slide.get("chart", {})
            categories, series = chart.get("categories", []), chart.get("series", [])
            if chart.get("type") not in {"column", "bar", "line"} or not 2 <= len(categories) <= 8 or not 1 <= len(series) <= 3:
                raise ValueError(f"Slide {n}: chart requires 2–8 categories and 1–3 series")
            if not chart.get("unit") or not slide.get("caption") or not evidence:
                raise ValueError(f"Slide {n}: chart unit, caption and evidence required")
            allowed = quoted_numbers(" ".join(e["quote"] for e in evidence)) | analysis_numbers
            for data in series:
                values = data.get("values", [])
                if not data.get("name") or len(values) != len(categories):
                    raise ValueError(f"Slide {n}: chart data dimensions do not match")
                for value in values:
                    if not isinstance(value, (int, float)) or not math.isfinite(value) or float(value) not in allowed:
                        raise ValueError(f"Slide {n}: chart value {value} lacks a quoted source value")
        if slide["layout"] == "table":
            rows = slide.get("table", [])
            if (not isinstance(rows, list) or not 2 <= len(rows) <= 6 or
                    not all(isinstance(row, list) for row in rows) or
                    not 2 <= len(rows[0]) <= 4 or any(len(row) != len(rows[0]) for row in rows)):
                raise ValueError(f"Slide {n}: table must be an array of rows (2–6 rows, 2–4 columns); do not use headers/rows object")
            if any(not isinstance(cell, str) or len(cell) > 65 for row in rows for cell in row) or not evidence:
                raise ValueError(f"Slide {n}: table cells too long or evidence missing")
        if slide["layout"] in {"process", "decision", "timeline", "key_metric", "comparison"}:
            visual = slide.get("visual", {})
            if visual.get("relationship", "sequence") not in {"sequence", "dimensions"}:
                raise ValueError(f"Slide {n}: unsupported visual relationship")
            key = {"process": "nodes", "decision": "nodes", "timeline": "items", "key_metric": "metrics", "comparison": "columns"}[slide["layout"]]
            items = visual.get(key)
            limit = {"process": (2, 7), "decision": (2, 5), "timeline": (2, 5), "key_metric": (1, 3), "comparison": (2, 2)}[slide["layout"]]
            if not isinstance(items, list) or not limit[0] <= len(items) <= limit[1] or not all(isinstance(item, dict) for item in items):
                raise ValueError(f"Slide {n}: {slide['layout']} needs {limit[0]}–{limit[1]} structured source-backed {key}")
            allowed = quoted_numbers(" ".join(e["quote"] for e in evidence)) | analysis_numbers
            for item in items:
                required = ("heading", "items") if slide["layout"] == "comparison" else ("label", "value") if slide["layout"] == "key_metric" else ("label",)
                if any(field not in item or item[field] in (None, "", []) for field in required):
                    raise ValueError(f"Slide {n}: visual item is missing {required}")
                if "items" in item and (not isinstance(item["items"], list) or not 1 <= len(item["items"]) <= 4 or any(not isinstance(text, str) or len(text) > 160 for text in item["items"])):
                    raise ValueError(f"Slide {n}: comparison requires 1–4 concise text items per column")
                for field in ("label", "detail", "value", "when", "heading"):
                    if field in item and (not isinstance(item[field], (str, int, float)) or len(str(item[field])) > 160):
                        raise ValueError(f"Slide {n}: visual {field} is invalid or too long")
                for field in ("value", "start", "end", "when"):
                    if field in item and not quoted_numbers(str(item[field])).issubset(allowed):
                        raise ValueError(f"Slide {n}: visual number in {field} lacks quoted source evidence")
                if item.get("locator") and item["locator"] not in {e["locator"] for e in evidence}:
                    raise ValueError(f"Slide {n}: visual item has an unquoted source locator")
        if contract.get("layout_version") == 2:
            from presentation.source_pages import slide_source_notes
            slide_source_notes(slide, sources)
    if contract.get("requires_charts") and not any(s["layout"] == "chart" for s in slides):
        raise ValueError("Order requires at least one editable chart with supplied numerical evidence")
    if analysis_packet is not None:
        from presentation.analysis import validate_slide_coverage
        validate_slide_coverage(spec, analysis_packet)
    from presentation.content_policy import validate_content_policy, visible_text
    validate_content_policy(spec, sources)
    ledger = None
    if contract.get("evidence_policy_version") == 2:
        from presentation.assurance import source_topics, validate_claim_ledger, validate_slide_claims
        ledger_path = workspace / "planning/presentation_claims.json"
        if not ledger_path.is_file():
            raise ValueError("Presentation claim ledger is missing; rebuild source evidence before slides")
        ledger = validate_claim_ledger(json.loads(ledger_path.read_text(encoding="utf-8")), sources, source_topics(sources))
        validate_slide_claims(spec, ledger, contract)
    claims = {claim["id"]: claim for claim in ledger["claims"]} if ledger else {}
    for n, slide in enumerate(slides, 1):
        if slide.get("section") == "title":
            # Cover identifiers can come from order metadata instead of the source.
            continue
        cited_text = " ".join(locators[e["locator"]] for e in slide.get("evidence", []))
        displayed = " ".join([*visible_text(slide), slide.get("notes", "")])
        displayed = re.sub(r"\b[A-Za-z]\d+:(?:block|table|page|image)-\d+\b", "", displayed, flags=re.I)
        order_context = " ".join(str(contract.get(key) or "") for key in ("title", "customer_notes"))
        calculated = {float(claims[identifier]["value"]) for identifier in slide.get("claim_ids", [])
                      if identifier in claims and claims[identifier]["kind"] == "calculated"}
        if analysis_packet is not None:
            from presentation.analysis import calculated_numbers
            calculated |= calculated_numbers(analysis_packet, slide.get("analysis_ids", []))
        unsupported = source_numbers(displayed) - source_numbers(cited_text + " " + order_context) - calculated
        if unsupported:
            raise ValueError(f"Slide {n}: numbers lack evidence from this slide's cited source passages: {sorted(unsupported)}")


def _style_paragraph(paragraph, text, size, color, font, rtl, bold=False):
    from pptx.util import Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.oxml.xmlchemy import OxmlElement
    paragraph.clear()
    paragraph.alignment = PP_ALIGN.RIGHT if rtl else PP_ALIGN.LEFT
    paragraph._p.get_or_add_pPr().set("rtl", "1" if rtl else "0")
    paragraph.space_after = Pt(14 if size >= 20 else 4)
    paragraph.line_spacing = 1.15
    # Give Latin phrases their own directional run. A single RTL run lets
    # PowerPoint attach adjacent Persian digits/punctuation to PET/NPV, or
    # reorder a wrapped English organization name. LRE/PDF plus an RLM after
    # the phrase return punctuation to the Persian paragraph's direction.
    parts = re.split(r"([A-Za-z][A-Za-z0-9_./+%\-]*(?:[ \t]+[A-Za-z][A-Za-z0-9_./+%\-]*)*)", text) if rtl else [text]
    for part in parts:
        if not part:
            continue
        latin = bool(rtl and re.match(r"^[A-Za-z]", part))
        run = paragraph.add_run()
        run.text = f"\u202a{part}\u202c\u200f" if latin else part
        run.font.name, run.font.size, run.font.bold = font, Pt(size), bold
        run.font.color.rgb = RGBColor.from_string(color)
        properties = run._r.get_or_add_rPr()
        properties.set("lang", "en-US" if latin or not rtl else "fa-IR")
        for tag in ("a:cs", "a:ea"):
            item = OxmlElement(tag)
            item.set("typeface", font)
            properties.append(item)


def build_pptx(spec: dict, contract: dict, sources: dict, target: Path, workspace: Path) -> None:
    if contract.get("layout_version") == 2:
        from presentation.compose import build_designed_pptx
        build_designed_pptx(spec, contract, sources, target, workspace)
        return
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.chart.data import CategoryChartData
    from PIL import Image
    from matplotlib.font_manager import FontProperties, findfont
    # A missing font is an installation issue, not permission to silently substitute.
    try:
        findfont(FontProperties(family=contract["font"]), fallback_to_default=False)
    except ValueError as exc:
        raise RuntimeError(f"Required font is not installed: {contract['font']}") from exc
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
    prs.core_properties.title = contract["title"]
    prs.core_properties.author = contract.get("student_name", "")
    rtl = contract.get("rtl", True)
    images = {image["id"]: image for source in sources["sources"] for image in source["images"]}
    navy, teal, ink = "102A43", "007D8A", "243B53"
    font = contract["font"]

    def box(slide, text, x, y, w, h, size, color=ink, bold=False, direction=None, references=False):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        shape.text_frame.word_wrap = True
        shape.text_frame.margin_left = shape.text_frame.margin_right = Inches(.04)
        shape.text_frame.margin_top = shape.text_frame.margin_bottom = Inches(.02)
        lines = re.split(r"\n|[؛;]\s*", text) if references else text.split("\n")
        if references:
            # Keep an English organization in its own paragraph rather than
            # wrapping it around unrelated Persian references. Filenames use
            # a separate LTR citation line, not an RTL body-text run.
            lines = [piece for line in lines for piece in re.split(r"(?<=:)\s+(?=[A-Za-z][A-Za-z0-9_.-]*\.(?:docx|pptx|pdf)\b)", line) if piece]
        for i, line in enumerate(lines):
            paragraph = shape.text_frame.paragraphs[0] if i == 0 else shape.text_frame.add_paragraph()
            paragraph_rtl = rtl if direction is None else direction
            paragraph_size = size
            if references:
                paragraph_rtl = bool(paragraph_rtl and any(ch.isalpha() and unicodedata.bidirectional(ch) in {"AL", "R"} for ch in line))
                paragraph_size = 16 if re.search(r"\.(docx|pptx|pdf)\b", line, re.I) else min(size, 20)
                if not paragraph_rtl:
                    line = re.sub(r"،\s*", ", ", line)
            _style_paragraph(paragraph, line, paragraph_size, color, font, paragraph_rtl, bold)
            if references:
                paragraph.space_after = Pt(6)
        return shape

    for index, data in enumerate(spec["slides"], 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string("FFFFFF")
        title_slide = data["layout"] == "title"
        if title_slide:
            slide.background.fill.fore_color.rgb = RGBColor.from_string(navy)
            box(slide, contract.get("university", ""), .85, .55, 11.6, .65, 19, "B6D7DB")
            box(slide, data["title"], .85, 1.65, 11.6, 2.0, 32, "FFFFFF", True)
            box(slide, "\n".join(data.get("points", [])), .85, 4.0, 11.6, 2.6, 21, "E0EEF0")
        else:
            stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(.12))
            stripe.fill.solid()
            stripe.fill.fore_color.rgb = RGBColor.from_string(teal)
            stripe.line.fill.background()
            box(slide, data["title"], .65, .4, 12.0, 1.25, contract["title_pt"], navy, True)
            points = "\n".join(data.get("points", []))
            layout = data["layout"]
            if layout == "statement":
                box(slide, points, .85, 2.0, 11.6, 4.55, contract["body_pt"], references=data["section"] == "references")
            elif layout == "two_column":
                values = data.get("points", [])
                midpoint = math.ceil(len(values) / 2)
                box(slide, "\n".join(values[:midpoint]), 6.9 if rtl else .85, 2.0, 5.5, 4.45, contract["body_pt"], references=data["section"] == "references")
                box(slide, "\n".join(values[midpoint:]), .85 if rtl else 6.9, 2.0, 5.5, 4.45, contract["body_pt"], references=data["section"] == "references")
            else:
                if layout == "image":
                    path = workspace / images[data["image_id"]]["path"]
                    with Image.open(path) as image:
                        width, height = image.size
                    scale = min(7.3 / width, 3.85 / height)
                    w, h = width * scale, height * scale
                    slide.shapes.add_picture(str(path), Inches(.8 + (7.3-w)/2), Inches(1.85 + (3.85-h)/2), Inches(w), Inches(h))
                elif layout == "chart":
                    chart_spec = data["chart"]
                    chart_data = CategoryChartData()
                    chart_data.categories = chart_spec["categories"]
                    for series in chart_spec["series"]:
                        chart_data.add_series(series["name"], series["values"])
                    chart_type = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS}[chart_spec["type"]]
                    chart = slide.shapes.add_chart(chart_type, Inches(.8), Inches(1.9), Inches(7.3), Inches(3.9), chart_data).chart
                    chart.has_title = False
                    chart.plots[0].vary_by_categories = False
                    for n, series in enumerate(chart.series):
                        series.format.line.color.rgb = RGBColor.from_string([teal, navy, "BF8300"][n])
                        if chart_spec["type"] != "line":
                            series.invert_if_negative = False
                            series.format.fill.solid()
                            series.format.fill.fore_color.rgb = RGBColor.from_string([teal, navy, "BF8300"][n])
                    chart.has_legend = len(chart_spec["series"]) > 1
                    if chart.has_legend:
                        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
                        chart.legend.font.name, chart.legend.font.size = font, Pt(15)
                    for axis in (chart.category_axis, chart.value_axis):
                        axis.tick_labels.font.name, axis.tick_labels.font.size = font, Pt(14)
                    chart.value_axis.has_title = True
                    chart.value_axis.axis_title.text_frame.text = chart_spec["unit"]
                    _style_paragraph(chart.value_axis.axis_title.text_frame.paragraphs[0], chart_spec["unit"], 14, ink, font, rtl)
                elif layout == "table":
                    rows = data["table"]
                    table = slide.shapes.add_table(len(rows), len(rows[0]), Inches(.8), Inches(1.9), Inches(7.3), Inches(3.9)).table
                    table._tbl.tblPr.set("rtl", "1" if rtl else "0")
                    for r, row in enumerate(rows):
                        for c, value in enumerate(row):
                            cell = table.cell(r, c)
                            cell.fill.solid()
                            cell.fill.fore_color.rgb = RGBColor.from_string(navy if r == 0 else "EFF5F7")
                            _style_paragraph(cell.text_frame.paragraphs[0], value, 17, "FFFFFF" if r == 0 else ink, font, rtl, r == 0)
                box(slide, points, 8.55, 1.95, 3.95, 3.9, 20)
                box(slide, data.get("caption", ""), .85, 5.98, 11.6, .62, 13, ink)
        refs = list(dict.fromkeys(e["locator"] for e in data.get("evidence", [])))
        if not title_slide:
            footer = " · ".join(refs)
            citation = box(slide, footer, .8, 6.85, 10.9, .48, max(13, contract["citation_pt"]), "526777", direction=False)
            citation.name = "Source citations"
        box(slide, str(index), 12.05, 6.92, .5, .3, 11, "B6D7DB" if title_slide else "526777")
        notes = data["notes"] + "\n\nSources:\n" + "\n".join(f"{e['locator']}: {e['quote']}" for e in data.get("evidence", []))
        slide.notes_slide.notes_text_frame.text = notes
    target.parent.mkdir(parents=True, exist_ok=True)
    prs.save(target)


def validate_package(path: Path, expected_count: int) -> dict:
    from pptx import Presentation
    prs = Presentation(path)
    errors = []
    if len(prs.slides) != expected_count:
        errors.append("PPTX slide count mismatch")
    for n, slide in enumerate(prs.slides, 1):
        if not slide.has_notes_slide or not slide.notes_slide.notes_text_frame.text.strip():
            errors.append(f"Slide {n}: notes missing")
        for shape in slide.shapes:
            if shape.left < 0 or shape.top < 0 or shape.left+shape.width > prs.slide_width+100 or shape.top+shape.height > prs.slide_height+100:
                errors.append(f"Slide {n}: object outside slide canvas")
    if errors:
        raise ValueError("; ".join(errors))
    return {"passed": True, "slides": len(prs.slides), "editable": True}


def render_powerpoint(pptx: Path, pdf: Path, pages: Path) -> None:
    """Render with an isolated Office document; never close the user's presentations."""
    import os
    import shutil
    import subprocess
    import tempfile
    import fitz
    pdf.parent.mkdir(parents=True, exist_ok=True)
    if os.name == "nt":
        # A child process bounds a hung COM server; no change to the active worker.
        subprocess.run([__import__('sys').executable, "-m", "presentation.render_office", str(pptx.resolve()), str(pdf.resolve())],
                       check=True, timeout=180, cwd=str(Path(__file__).resolve().parents[1]))
    else:
        office = shutil.which("libreoffice") or shutil.which("soffice")
        if not office:
            raise RuntimeError("Presentation renderer temporarily unavailable: install PowerPoint on Windows or LibreOffice on Linux.")
        with tempfile.TemporaryDirectory() as profile:
            subprocess.run([office, f"-env:UserInstallation={Path(profile).as_uri()}", "--headless", "--convert-to", "pdf", "--outdir", str(pdf.parent), str(pptx)], check=True, timeout=180)
        exported = pdf.parent / pptx.with_suffix(".pdf").name
        if exported != pdf:
            exported.replace(pdf)
    pages.mkdir(parents=True, exist_ok=True)
    with fitz.open(pdf) as document:
        for n, page in enumerate(document, 1):
            page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5)).save(pages / f"slide-{n:03}.png")
