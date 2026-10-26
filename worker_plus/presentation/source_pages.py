"""Truthful source page ranges for speaker notes; never estimate pages from length."""
from __future__ import annotations

import hashlib
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys


def _bookmark(document, name, first, last, number):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    begin, end = OxmlElement("w:bookmarkStart"), OxmlElement("w:bookmarkEnd")
    begin.set(qn("w:id"), str(number)); begin.set(qn("w:name"), name)
    end.set(qn("w:id"), str(number))
    # Keep paragraph properties first; bookmarks do not affect pagination.
    first.insert(1 if len(first) and first[0].tag == qn("w:pPr") else 0, begin)
    last.append(end)


def paginate_source(source: dict, original: Path, destination: Path) -> dict:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt
    import fitz

    if hashlib.sha256(original.read_bytes()).hexdigest() != source["sha256"]:
        raise ValueError("Source changed before pagination")
    destination.mkdir(parents=True, exist_ok=True)
    suffix = original.suffix.lower()
    pdf = destination / f"{source['id']}.pdf"
    if suffix == ".pdf":
        shutil.copy2(original, pdf)
        with fitz.open(pdf) as document:
            count = len(document)
        method = "original_pdf"
    elif suffix == ".pptx":
        # A slide is a real page in the original presentation, not a guessed
        # Word page. Its number is already present in the source locator.
        from pptx import Presentation
        count = len(Presentation(original).slides)
        method = "original_slides"
    else:
        if os.name != "nt":
            raise RuntimeError("Source pagination requires Microsoft Word on Windows; no page numbers were invented")
        document = Document(original) if suffix == ".docx" else Document()
        if suffix != ".docx":
            document.styles["Normal"].font.name = "Tahoma"
            document.styles["Normal"].font.size = Pt(12)
            for line in original.read_text(encoding="utf-8-sig").splitlines():
                paragraph = document.add_paragraph(line)
                if any("\u0600" <= char <= "\u06ff" for char in line):
                    paragraph._p.get_or_add_pPr().append(__import__("docx.oxml", fromlist=["OxmlElement"]).OxmlElement("w:bidi"))
        bookmark_ids = {}
        elements = list(document.element.body)
        anchors = list(dict.fromkeys([s["id"] for s in source["segments"]] +
                                    [image["source_locator"] for image in source["images"] if image.get("source_locator")]))
        for number, identifier in enumerate(anchors, 1):
            locator = identifier.split(":", 1)[1]
            position = int(locator.rsplit("-", 1)[1])
            element = elements[position - 1]
            if element.tag == qn("w:tbl"):
                paragraphs = list(element.iter(qn("w:p")))
                first, last = paragraphs[0], paragraphs[-1]
            else:
                first = last = element
            name = f"WPSource_{source['id']}_{number}"
            _bookmark(document, name, first, last, 50000 + number)
            bookmark_ids[identifier] = name
        annotated = destination / f"{source['id']}-pagination.docx"
        result = destination / f"{source['id']}-pagination.json"
        document.save(annotated)
        subprocess.run([sys.executable, "-m", "presentation.source_office", str(annotated), str(pdf), str(result)],
                       cwd=str(Path(__file__).resolve().parents[1]), check=True, timeout=240)
        pagination = json.loads(result.read_text(encoding="utf-8"))
        for segment in source["segments"]:
            pages = pagination["bookmarks"].get(bookmark_ids[segment["id"]])
            if not pages:
                raise ValueError(f"Source page mapping missing: {segment['id']}")
            segment["pages"] = pages
        for image in source["images"]:
            if image.get("source_locator") in bookmark_ids:
                image["pages"] = pagination["bookmarks"].get(bookmark_ids[image["source_locator"]])
        count = pagination["page_count"]
        method = "word_physical_pages" if suffix == ".docx" else "paginated_text_snapshot"
    for segment in source["segments"]:
        if not segment.get("pages") or any(type(page) is not int or not 1 <= page <= count for page in segment["pages"]):
            raise ValueError(f"No verified page range for {segment['id']}")
    lookup = {s["id"]: s["pages"] for s in source["segments"]}
    for image in source["images"]:
        if image.get("source_locator") in lookup:
            image["pages"] = lookup[image["source_locator"]]
    source["pagination"] = {"method": method, "page_count": count, "source_sha256": source["sha256"],
                            "snapshot": str(pdf) if pdf.exists() else None, "numbering": "physical_from_file_start"}
    return source


def merge_pages(pages):
    result = []
    for page in sorted(set(pages)):
        if result and page == result[-1][1] + 1:
            result[-1][1] = page
        else:
            result.append([page, page])
    return result


def slide_source_notes(slide: dict, sources: dict, *, cover=False) -> list[dict]:
    selected = {e["locator"] for e in slide.get("evidence", [])}
    selected.update(slide.get("source_locators", []))
    image_id = slide.get("image_id")
    notes = []
    for index, source in enumerate(sources["sources"], 1):
        segments = [s for s in source["segments"] if s["id"] in selected]
        images = [image for image in source["images"] if image["id"] == image_id]
        if not selected and not images:
            # Orientation for non-factual cover/agenda/question slides, clearly
            # labelled as such rather than pretending the source proves them.
            segments = source["segments"][:1]
        entries = segments + images
        if not entries:
            continue
        if any(not entry.get("pages") for entry in entries):
            raise ValueError(f"Verified source pages are required for speaker notes: {source['name']}")
        pages = [page for entry in entries for page in entry["pages"]]
        notes.append({"source_id": source["id"], "source_number": index, "file_name": source.get("original_name", source["name"]),
                      "ranges": merge_pages(pages), "locators": [s["id"] for s in entries],
                      "pagination": source.get("pagination", {}), "purpose": "orientation" if not selected and not images else "slide_evidence"})
    if selected - {s["id"] for source in sources["sources"] for s in source["segments"]}:
        raise ValueError("Speaker notes contain an unknown source locator")
    if not notes:
        raise ValueError("Every slide requires a source preparation note")
    return notes


def format_source_notes(items, rtl=True):
    lines = ["راهنمای مطالعه و آمادگی ارائه" if rtl else "Source reading and presentation preparation"]
    for item in items:
        ranges = ", ".join(str(a) if a == b else f"{a}–{b}" for a, b in item["ranges"])
        is_slides = item["pagination"].get("method") == "original_slides"
        label = ("اسلایدها" if is_slides else "صفحات") if rtl else ("slides" if is_slides else "pages")
        kind = " (برای معرفی منبع)" if item["purpose"] == "orientation" and rtl else ""
        lines.append(f"[{item['source_number']}] {item['file_name']} · {label} {ranges}{kind}")
        if item["pagination"].get("method") == "paginated_text_snapshot":
            lines.append("شماره‌ها مربوط به نسخهٔ صفحه‌بندی‌شدهٔ فایل متنی‌اند." if rtl else "Page numbers refer to the paginated text snapshot.")
    return "\n".join(lines)
