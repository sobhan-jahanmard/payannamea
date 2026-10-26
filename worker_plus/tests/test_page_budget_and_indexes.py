import importlib
import json
import unittest
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import patch

from common.page_budget import page_budget
from common.profile_support import DEFAULT_DOCUMENT_RULES
from common.runner import STEPS, Services, recovery_start


class PageBudgetAndIndexesTests(unittest.TestCase):
    def test_twenty_page_order_has_an_upper_bound(self) -> None:
        budget = page_budget({"quantity_type": "pages", "quantity_value": 20})
        self.assertEqual({
            "minimum": 20, "maximum": 24, "word_maximum": 23,
            "minimum_content_words": 2793, "minimum_persian_words": 2513,
            "suggested_content_words": 2940, "maximum_content_words": 3381,
            "reserved_nonbody_pages": 6, "estimated_visuals": 2,
            "estimated_tables": 1, "reserved_visual_pages": 1,
            "reserved_table_pages": 1, "estimated_text_pages": 12,
            "words_per_text_page": 245,
        }, budget)
        hundred = page_budget({"quantity_type": "pages", "quantity_value": 100})
        self.assertEqual((100, 110, 85, 20825, 23949), (
            hundred["minimum"], hundred["maximum"], hundred["estimated_text_pages"],
            hundred["suggested_content_words"], hundred["maximum_content_words"],
        ))
        self.assertLess(page_budget({"quantity_type": "pages", "quantity_value": 100}, visual_count=12)["suggested_content_words"], hundred["suggested_content_words"])
        denser_rules = {"font": {"body_pt": 12}, "paragraph": {"line_spacing": 1.2}}
        self.assertGreater(page_budget({"quantity_type": "pages", "quantity_value": 100}, rules=denser_rules)["suggested_content_words"], hundred["suggested_content_words"])
        self.assertIsNone(page_budget({"quantity_type": "words", "quantity_value": 20}))
        self.assertEqual("07_generate_content", STEPS[recovery_start("10_validate_and_publish", RuntimeError("Rendered Word package missed the 20–23 page window"))])

    def test_oversized_first_draft_is_compressed_before_word_packaging(self) -> None:
        stage = importlib.import_module("common.steps.07_generate_content")
        with TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "final").mkdir()
            prompts = []
            source = root / "final/deliverable_source.md"

            def run_codex(prompt, report):
                prompts.append(prompt)
                if report == source:
                    source.write_text("# چکیده\n" + "تحلیل " * 5000, encoding="utf-8")
                else:
                    source.write_text("# چکیده\n" + "تحلیل " * 2940, encoding="utf-8")
                    report.write_text("compressed", encoding="utf-8")

            services = SimpleNamespace(
                workspace=root, profile=SimpleNamespace(DISPLAY_NAME="گزارش"),
                recovery_feedback=None, run_codex=run_codex,
            )
            context = {"order": {"title": "طرح", "quantity_type": "pages", "quantity_value": 20}, "artifacts": {
                "research_evidence": "research/evidence.md", "research_analysis": "research/analysis.md",
                "plan": "planning/outline.md",
            }}
            stage.run(context, services)
            self.assertEqual(2, len(prompts))
            self.assertIn("2793 تا 3381", prompts[0])
            self.assertIn("پیش از صفحه‌آرایی Word", prompts[1])
            self.assertEqual(2941, stage.source_word_count(source.read_text(encoding="utf-8")))
            report = json.loads((root / "reports/stage_checks/draft_content_budget.json").read_text(encoding="utf-8"))
            self.assertEqual(12, report["estimated_text_pages"])

    def test_hundred_page_order_extends_in_bounded_batches(self) -> None:
        stage = importlib.import_module("common.steps.07_generate_content")
        with TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "final").mkdir()
            source = root / "final/deliverable_source.md"
            prompts = []

            def run_codex(prompt, report):
                prompts.append(prompt)
                if report == source:
                    source.write_text("# چکیده\n" + "تحلیل " * 5000, encoding="utf-8")
                else:
                    source.write_text(source.read_text(encoding="utf-8") + "تحلیل " * 3000, encoding="utf-8")
                    report.write_text("extended", encoding="utf-8")

            services = SimpleNamespace(workspace=root, profile=SimpleNamespace(DISPLAY_NAME="گزارش"), recovery_feedback=None, run_codex=run_codex)
            context = {"order": {"title": "طرح", "quantity_type": "pages", "quantity_value": 100}, "artifacts": {
                "research_evidence": "research/evidence.md", "research_analysis": "research/analysis.md", "plan": "planning/outline.md",
            }}
            stage.run(context, services)
            self.assertIn("100 تا 110 صفحه", prompts[0])
            self.assertEqual(6, len(prompts))
            self.assertTrue(all("حداکثر 3000 واژه" in prompt for prompt in prompts[1:-1]))
            self.assertIn("حداکثر 2932 واژه", prompts[-1])
            self.assertLessEqual(stage.source_word_count(source.read_text(encoding="utf-8")), page_budget(context["order"])["maximum_content_words"])

    def test_source_audit_does_not_force_six_thousand_words_for_twenty_pages(self) -> None:
        stage = importlib.import_module("common.steps.08b_audit_source_compliance")
        with TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "source.md").write_text("واژه " * 3000 + "\n# منابع\n", encoding="utf-8")
            (root / "rules.md").write_text("قواعد منابع", encoding="utf-8")
            context = {"order": {"quantity_type": "pages", "quantity_value": 20}, "artifacts": {"source": "source.md", "resolved_source_rules": "rules.md"}}
            passed, checks = stage.evaluate(context, SimpleNamespace(workspace=root))
            self.assertTrue(passed, checks)

    def test_page_feedback_compresses_and_rechecks_content(self) -> None:
        stage = importlib.import_module("common.steps.10_validate_and_publish")
        with TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "drafts").mkdir()
            source = root / "source.md"
            source.write_text("original", encoding="utf-8")
            docx = root / "output.docx"
            docx.write_bytes(b"fixture")
            counts = iter([66, 22])
            prompts = []
            review_calls = []

            def run_codex(prompt, report):
                prompts.append(prompt)
                source.write_text("compact", encoding="utf-8")
                report.write_text("done", encoding="utf-8")

            services = SimpleNamespace(workspace=root, count_docx_pages=lambda _: next(counts), run_codex=run_codex)
            context = {"order": {"quantity_type": "pages", "quantity_value": 20}, "artifacts": {"order_execution_contract": "contract.md"}}
            with patch.object(stage, "repair_layout") as replay, patch.object(stage.importlib, "import_module", return_value=SimpleNamespace(run=lambda *_: review_calls.append(True))):
                self.assertEqual(22, stage.complete_page_target(context, services, source, docx, {}))
            self.assertEqual(1, len(prompts))
            self.assertIn("66", prompts[0])
            self.assertEqual([True, True], review_calls)
            replay.assert_called_once()
            report = json.loads((root / "reports/stage_checks/page_completion.json").read_text(encoding="utf-8"))
            self.assertTrue(report["passed"])
            self.assertEqual(["compress", "pass"], [item["action"] for item in report["attempts"]])

    def test_page_feedback_expands_in_small_batch(self) -> None:
        stage = importlib.import_module("common.steps.10_validate_and_publish")
        with TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "drafts").mkdir()
            source = root / "source.md"
            source.write_text("short", encoding="utf-8")
            docx = root / "output.docx"
            docx.write_bytes(b"fixture")
            counts = iter([19, 21])
            prompts = []

            def run_codex(prompt, report):
                prompts.append(prompt)
                source.write_text("expanded", encoding="utf-8")
                report.write_text("done", encoding="utf-8")

            services = SimpleNamespace(workspace=root, count_docx_pages=lambda _: next(counts), run_codex=run_codex)
            context = {"order": {"quantity_type": "pages", "quantity_value": 20}, "artifacts": {"order_execution_contract": "contract.md"}}
            with patch.object(stage, "repair_layout") as replay, patch.object(stage.importlib, "import_module", return_value=SimpleNamespace(run=lambda *_: None)):
                self.assertEqual(21, stage.complete_page_target(context, services, source, docx, {}))
            self.assertEqual(1, len(prompts))
            self.assertIn("245", prompts[0])
            replay.assert_called_once()

    def test_indexes_use_separate_title_and_page_cells_with_bookmarks(self) -> None:
        from PIL import Image
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn

        with TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "figures").mkdir()
            Image.new("RGB", (20, 20), "white").save(root / "figures" / "chart.png")
            source = root / "source.md"
            source.write_text(
                "# چکیده\nمتن چکیده.\n# فهرست مطالب\n- فصل نادرست و قدیمی\n"
                "# فصل ۱: بررسی\n## ۱-۱ داده\nجدول ۱. داده‌های واقعی — منبع: گزارش رسمی\n"
                "| مقدار | واحد |\n| --- | --- |\n| ۲ | تن |\n"
                "![شکل ۱. نمودار داده‌ها؛ منبع: گزارش رسمی](figures/chart.png)\n"
                "# منابع\n[1] مرجع.\n", encoding="utf-8",
            )
            services = Services.__new__(Services)
            services.release_docx_lock = lambda _: None
            output = root / "output.docx"
            services.write_docx(source, output, "عنوان", DEFAULT_DOCUMENT_RULES, {"title": "عنوان", "student_name": "دانشجو"})
            document = Document(output)
            self.assertGreaterEqual(len(document.tables), 4)
            toc, figures, tables = document.tables[:3]
            self.assertEqual(WD_ALIGN_PARAGRAPH.LEFT, toc.cell(0, 0).paragraphs[0].alignment)
            self.assertEqual(WD_ALIGN_PARAGRAPH.RIGHT, toc.cell(0, 1).paragraphs[0].alignment)
            self.assertNotIn("قدیمی", " ".join(cell.text for row in toc.rows for cell in row.cells))
            self.assertIn("شکل ۱", figures.cell(0, 0).text)
            self.assertNotIn("منبع", figures.cell(0, 0).text)
            self.assertIn("جدول ۱", tables.cell(0, 0).text)
            self.assertNotIn("منبع", tables.cell(0, 0).text)
            anchors = {item.get(qn("w:name")) for item in document._element.iter(qn("w:bookmarkStart"))}
            links = {item.get(qn("w:anchor")) for item in document._element.iter(qn("w:hyperlink"))}
            self.assertTrue(links <= anchors)
            self.assertTrue(any(anchor.startswith("figure_") for anchor in links))
            self.assertTrue(any(anchor.startswith("table_") for anchor in links))
            self.assertEqual(len(links), len([item for item in document._element.iter(qn("w:instrText")) if "PAGEREF" in (item.text or "")]))

    def test_course_research_cover_and_sections_use_report_format(self) -> None:
        from docx import Document
        from university_research.profile import resolve_rules

        with TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "source.md"
            source.write_text("# چکیده\nخلاصه.\n# فهرست مطالب\n# ۱ مقدمه\nمتن.\n# ۲ یافته‌ها\nنتیجه.\n# منابع\n[1] منبع.\n", encoding="utf-8")
            output = root / "report.docx"
            order = {"order_type": "تحقیق دانشگاهی", "title": "عنوان تحقیق", "university": "دانشگاه نمونه", "student_name": "دانشجو", "course_name": "تحقیق در عملیات", "instructor_name": "استاد درس"}
            services = Services.__new__(Services)
            services.release_docx_lock = lambda _: None
            rules = resolve_rules([])
            services.write_docx(source, output, order["title"], rules, order)
            services.polish_cover(output, rules, order)
            document = Document(output)
            cover = [p.text for p in document.paragraphs if p.style.name == "Title"]
            self.assertIn("گزارش تحقیق درس تحقیق در عملیات", cover)
            self.assertIn("استاد درس: استاد درس", cover)
            self.assertFalse(any("دریافت درجه" in line or "استاد راهنما" in line for line in cover))
            sections = [p for p in document.paragraphs if p.style.name == "Heading 1"]
            self.assertTrue(all(not p.paragraph_format.page_break_before for p in sections[1:]))


if __name__ == "__main__":
    unittest.main()
