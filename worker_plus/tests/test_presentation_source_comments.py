import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
import zipfile

from pptx import Presentation
from presentation.source_comments import add_source_comments, validate_source_comments, INTERNAL_REFERENCE


class SourceCommentsTests(unittest.TestCase):
    def test_internal_reference_detection_does_not_ban_scientific_sample_names(self):
        self.assertIsNone(INTERNAL_REFERENCE.search('Sample S1 has the measured response'))
        for label in ['Source S1 | pp. 2', 'S1:block-12', 'منبع S2']:
            self.assertIsNotNone(INTERNAL_REFERENCE.search(label))

    def fixture(self, root):
        path = root / 'comments.pptx'
        prs = Presentation(); slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.notes_slide.notes_text_frame.text = 'Speaker narration'
        prs.save(path)
        texts = ['study.docx · pages 2, 5–6\nS1:block-12: exact supplied passage']
        add_source_comments(path, texts)
        return path, texts

    def test_comment_relationships_are_native_idempotent_and_not_slide_text(self):
        with tempfile.TemporaryDirectory() as directory:
            path, texts = self.fixture(Path(directory))
            add_source_comments(path, texts)
            self.assertTrue(validate_source_comments(path, texts)['passed'])
            with zipfile.ZipFile(path) as z:
                self.assertNotIn(b'S1:block-12', z.read('ppt/slides/slide1.xml'))
                self.assertIn(b'S1:block-12', z.read('ppt/comments/comment1.xml'))
                self.assertIn(b'commentAuthors+xml', z.read('[Content_Types].xml'))
                self.assertIn(b'/comments', z.read('ppt/slides/_rels/slide1.xml.rels'))
            self.assertIn('Speaker narration', Presentation(path).slides[0].notes_slide.notes_text_frame.text)
            with self.assertRaisesRegex(ValueError, 'stale'):
                validate_source_comments(path, ['A changed expected passage'])

    @unittest.skipUnless(os.name == 'nt', 'PowerPoint COM requires Windows')
    def test_actual_powerpoint_reads_the_comment_text(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); path, texts = self.fixture(root)
            script = '''import json,sys,pythoncom,win32com.client
pythoncom.CoInitialize()
app=doc=None
owns=False
try:
 try: app=win32com.client.GetActiveObject('PowerPoint.Application')
 except pythoncom.com_error:
  app=win32com.client.DispatchEx('PowerPoint.Application');owns=True
 doc=app.Presentations.Open(sys.argv[1],ReadOnly=True,Untitled=False,WithWindow=False)
 rows=[str(doc.Slides(1).Comments(i).Text) for i in range(1,doc.Slides(1).Comments.Count+1)]
 open(sys.argv[2],'w',encoding='utf-8').write(json.dumps(rows,ensure_ascii=False))
finally:
 if doc is not None: doc.Close()
 if app is not None and owns: app.Quit()
 pythoncom.CoUninitialize()
'''
            child = root / 'native_check.py'; child.write_text(script, encoding='utf-8')
            subprocess.run([sys.executable, str(child), str(path), str(root/'result.json')], check=True, timeout=60)
            self.assertEqual(texts, json.loads((root/'result.json').read_text(encoding='utf-8')))


if __name__ == '__main__': unittest.main()
