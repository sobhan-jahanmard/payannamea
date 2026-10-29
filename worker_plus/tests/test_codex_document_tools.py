"""Document lifecycle checks with local fixtures. No research, model or backend calls."""
import copy
from contextlib import redirect_stdout
from io import StringIO
import json
from pathlib import Path
import tempfile
import unittest
from types import SimpleNamespace
from unittest.mock import patch

from codex_order_tools import OrderTools, SKILLS, hashes, main, signature
from codex_document_tools import DocumentTools, BASE, word_count
from common.helpers import write_json, write_text


TYPES = [kind for kind in SKILLS if kind != "ارائه و پاورپوینت"]


def fixture(root, kind="تحقیق دانشگاهی"):
    order = {"id": "document-test", "order_type": kind, "title": "سنجش شواهد",
             "faculty": "آزمایشی", "advisor_name": "استاد نمونه", "abstract": "شرح مسئله",
             "instructor_name": "استاد درس", "course_name": "روش تحقیق", "quantity_type": "words",
             "quantity_value": 5, "language": "persian", "files": []}
    snapshot = root / "order.json"
    write_json(snapshot, order)
    tool = DocumentTools(root / "work")
    tool.fetch(order_file=snapshot)
    tool.prepare()
    tool.extract()
    write_text(tool.workspace / "planning/resolved_source_rules.md", "قواعد این سفارش؛ بدون شیوه‌نامهٔ آپلودی.\n")
    required = json.loads((tool.workspace / BASE[1]).read_text(encoding="utf-8"))["deliverable_contract"]["deliverables"]
    evidence = {"deliverables": [{"id": item["id"], "status": "direct",
        "sources": [{"url": "https://example.test/evidence", "title": "Fixture only", "publisher": "Fixture",
                     "geography": "international", "retrieved_at": "2026-10-28", "extracted_fact": "Fixture evidence"}],
        "method_executed": "Fixture synthesis", "result": "Fixture conclusion", "limitations": "Fixture only",
        "calculation_or_comparison": "Fixture comparison"} for item in required]}
    write_json(root / "evidence.json", evidence)
    write_text(tool.workspace / "research/analysis.md", "Local fixture, not real academic research.\n")
    tool.validate("evidence", root / "evidence.json")
    plan = {"sections": [{"id": "intro", "title": "مقدمه", "evidence_ids": [required[0]["id"]]}], "companions": []}
    write_json(tool.workspace / "planning/document_plan.json", plan)
    write_json(tool.workspace / "planning/visual_plan.json", {"visuals": []})
    write_json(tool.workspace / "final/figures/manifest.json", {"figures": []})
    write_text(root / "candidate.md", "# مقدمه\nاین متن صرفاً برای بررسی ابزار نوشته شده است [1].\n# منابع\n[1] Fixture, not a real reference.\n")
    tool.validate("document", root / "candidate.md")
    return tool, snapshot


def audit(tool, path, kind="content", status="PASS", issues=None):
    inputs = tool.review_inputs(kind)["inputs"]
    details = {"reviewed_sections": ["intro"], "issues": issues or []}
    write_text(path, f"STATUS: {status}\nINPUTS_SHA256: " + json.dumps(inputs) +
               "\nISSUES_JSON: " + json.dumps(details) + "\nFixture validation only.\n")


class DocumentToolsTest(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.tool, self.snapshot = fixture(self.root)

    def test_every_frontend_type_has_executable_route_and_installed_skill_source(self):
        repo = Path(__file__).resolve().parents[2]
        source = (repo / "frontend/src/lib/order-options.ts").read_text(encoding="utf-8")
        import re
        values = re.findall(r'"([^"]+)"', source.split("export const orderTypeOptions = [", 1)[1].split("];", 1)[0])
        self.assertEqual(set(values), set(SKILLS))
        for name in SKILLS.values():
            self.assertTrue((repo / "codex_skills" / name / "SKILL.md").is_file())

    def test_all_document_profiles_complete_intake_evidence_and_document_validation(self):
        for n, kind in enumerate(TYPES):
            root = self.root / str(n)
            root.mkdir()
            with self.subTest(kind=kind), patch("codex_order_tools.api.claim_oldest") as claim:
                tool, snapshot = fixture(root, kind)
                self.assertTrue(tool.current("document"))
                self.assertEqual(tool.fetch(order_file=snapshot)["skill"], SKILLS[kind])
                self.assertEqual(tool.state()["source_policy"], "supplied_and_verified_research")
                claim.assert_not_called()

    def test_cli_dispatches_from_saved_order_type(self):
        stream = StringIO()
        with redirect_stdout(stream):
            result = main(["--workspace", str(self.tool.workspace), "progress", "--step", "5"])
        self.assertEqual(result, 0)
        self.assertEqual(json.loads(stream.getvalue())["result"]["label"], "طراحی ساختار و روش")

    def test_original_data_tampering_blocks_document_actions(self):
        with (self.tool.workspace / "customer_input.json").open("a", encoding="utf-8") as handle:
            handle.write(" ")
        with self.assertRaisesRegex(ValueError, "Original order inputs changed"):
            self.tool.validate("document")

    def test_rejected_candidate_preserves_last_valid_source(self):
        before = hashes(self.tool.workspace, ["final/deliverable_source.md"])
        write_text(self.root / "invalid.md", "No numbered references")
        with self.assertRaisesRegex(ValueError, "Missing planned sections"):
            self.tool.validate("document", self.root / "invalid.md")
        self.assertEqual(before, hashes(self.tool.workspace, before))

    def test_bad_citation_and_word_target_are_rejected(self):
        text = (self.root / "candidate.md").read_text(encoding="utf-8").replace("[1].", "[9].")
        write_text(self.root / "wrong.md", text)
        with self.assertRaisesRegex(ValueError, "Unmatched"):
            self.tool.validate("document", self.root / "wrong.md")
        write_text(self.root / "wrong.md", "# مقدمه\nکم [1]\n# منابع\n[1] Fixture\n")
        with self.assertRaisesRegex(ValueError, "word count below"):
            self.tool.validate("document", self.root / "wrong.md")

    def test_stale_and_failed_reviews_cannot_preserve_approval(self):
        path = self.root / "audit.md"
        audit(self.tool, path)
        self.tool.review("content", path)
        self.assertTrue(self.tool.current("review_content"))
        audit(self.tool, path, status="FAIL", issues=[{"section": "intro", "message": "Fixture issue"}])
        with self.assertRaisesRegex(ValueError, "Review failed"):
            self.tool.review("content", path)
        self.assertFalse(self.tool.current("review_content"))
        audit(self.tool, path)
        write_text(self.tool.workspace / "research/analysis.md", "Changed analysis\n")
        with self.assertRaisesRegex(ValueError, "Missing or stale|snapshot missing or stale"):
            self.tool.review("content", path)

    def test_uploaded_guidelines_and_sources_stay_distinct(self):
        root = self.root / "tagged"
        root.mkdir()
        source = root / "source.txt"
        source.write_text("Actual local fixture content", encoding="utf-8")
        order = json.loads(self.snapshot.read_text(encoding="utf-8"))
        order["files"] = [{"id": "S1", "file_type": "information_source", "original_name": source.name,
                           "size_bytes": source.stat().st_size, "local_path": str(source)}]
        write_json(root / "order.json", order)
        tool = DocumentTools(root / "work")
        tool.fetch(order_file=root / "order.json")
        tool.prepare()
        tool.extract()
        index = json.loads((tool.workspace / "extracted/document_sources.json").read_text(encoding="utf-8"))
        self.assertEqual(index["sources"][0]["role"], "information_source")
        self.assertIn("S1:line-1", index["sources"][0]["segments"][0]["id"])
        self.assertEqual(tool.context()["artifacts"]["input_files"], [])
        write_text(tool.workspace / "planning/resolved_source_rules.md", "Supplied-only fixture\n")
        write_text(tool.workspace / "research/analysis.md", "Private source synthesis fixture\n")
        evidence = json.loads((self.root / "evidence.json").read_text(encoding="utf-8"))
        for entry in evidence["deliverables"]:
            entry["sources"] = []
            entry["uploaded_evidence"] = [{"locator": "S1:line-1", "quote": "Actual local fixture content"}]
        write_json(root / "evidence.json", evidence)
        tool.validate("evidence", root / "evidence.json")
        self.assertTrue(tool.current("evidence"))
        before = hashes(tool.workspace, ["research/evidence.json"])
        evidence["deliverables"][0]["uploaded_evidence"][0]["quote"] = "Invented quotation"
        write_json(root / "evidence.json", evidence)
        with self.assertRaisesRegex(ValueError, "unsupported quote"):
            tool.validate("evidence", root / "evidence.json")
        self.assertEqual(before, hashes(tool.workspace, before))

    def test_changed_academic_fields_change_document_signature(self):
        order = json.loads(self.snapshot.read_text(encoding="utf-8"))
        for key in ("methodology", "references", "review_notes", "advisor_name"):
            modified = copy.deepcopy(order)
            modified[key] = "Changed fixture"
            self.assertNotEqual(signature(order), signature(modified))

    def test_companion_edits_invalidate_document_and_reviews(self):
        path = self.tool.workspace / "planning/document_plan.json"
        plan = json.loads(path.read_text(encoding="utf-8"))
        plan["companions"] = [{"title": "گزارش", "source": "final/report_1.md"}]
        write_json(path, plan)
        write_text(self.tool.workspace / "final/report_1.md", "First companion\n")
        self.tool.validate("document")
        self.assertTrue(self.tool.current("document"))
        write_text(self.tool.workspace / "final/report_1.md", "Changed companion\n")
        self.assertFalse(self.tool.current("document"))

    def test_unreviewed_documents_cannot_build_or_publish(self):
        with self.assertRaisesRegex(ValueError, "review_content"):
            self.tool.build()
        with patch("codex_order_tools.api.claim_oldest") as claim:
            with self.assertRaisesRegex(ValueError, "review_content"):
                self.tool.publish()
            claim.assert_not_called()

    def test_document_service_never_starts_model(self):
        with patch("common.runner.Services.automation_word_pids", return_value=set()):
            service = self.tool.services()
        with self.assertRaisesRegex(RuntimeError, "model subprocesses are disabled"):
            service.run_codex("Any prompt", self.root / "never-written.md")

    def test_native_cover_uses_actual_document_type(self):
        from common.runner import Services
        from common.profile_support import DEFAULT_DOCUMENT_RULES
        from docx import Document
        for kind in TYPES:
            service = Services.__new__(Services)
            service.release_docx_lock = lambda _: None
            order = {"order_type": kind, "title": "Fixture", "student_name": "دانشجو", "course_name": "درس"}
            output = self.root / "cover.docx"
            service.write_docx(self.root / "candidate.md", output, "Fixture", DEFAULT_DOCUMENT_RULES, order)
            service.polish_cover(output, DEFAULT_DOCUMENT_RULES, order)
            cover = "\n".join(p.text for p in Document(output).paragraphs if p.style.name == "Title")
            self.assertNotIn("پروژه طراحی", cover)
            if kind == "پروپوزال پایان‌نامه":
                self.assertIn("پروپوزال پایان‌نامه", cover)
                self.assertNotIn("دریافت درجه", cover)
            elif kind == "پایان‌نامه کارشناسی ارشد":
                self.assertIn("کارشناسی ارشد", cover)

    def test_unicode_words_exclude_numeric_citations_and_image_paths(self):
        self.assertEqual(word_count("سلام دنیا [۱۲] ![تصویر](figures/1.png)"), 2)

    def rendered_fixture(self, sample_pages=2):
        """Exercise orchestration with valid files; only the Office boundary is stubbed."""
        from docx import Document
        import fitz
        tool = self.tool
        for kind in ("content", "sources"):
            report = self.root / f"{kind}.md"
            audit(tool, report, kind)
            tool.review(kind, report)
        native = tool.workspace / "final/final.docx"
        Document().save(native)
        tool.record("build", [*tool.content_paths(), "final/final.docx"])
        def export(docx, pdf, font):
            pages = sample_pages if docx.name == "sample.docx" else 4
            with fitz.open() as document:
                for n in range(pages):
                    page = document.new_page()
                    page.insert_text((40, 40), f"Local fixture page {n + 1}")
                document.save(pdf)
        def sample(docx, output):
            Document().save(output)
            return 4, sample_pages
        service = SimpleNamespace(export_pdf=export, create_first_half_sample=sample)
        with patch.object(tool, "services", return_value=service):
            result = tool.render()
        visual = self.root / "visual.md"
        details = {"reviewed_images": list(result["images"]), "issues": []}
        write_text(visual, "STATUS: PASS\nINPUTS_SHA256: " + json.dumps(tool.review_inputs("visual")["inputs"]) +
                   "\nISSUES_JSON: " + json.dumps(details) + "\nFixture page rendering only.\n")
        tool.review("visual", visual)
        return result

    def test_render_and_finalize_cover_full_and_sample_page_images(self):
        result = self.rendered_fixture()
        self.assertEqual(len(result["images"]), 6)
        report = self.tool.finalize()
        self.assertTrue(report["passed"])
        self.assertEqual(set(report["outputs"]), set(self.tool.outputs))
        image = self.tool.workspace / next(iter(result["images"]))
        image.write_bytes(b"Changed image after review")
        self.assertFalse(self.tool.current("render"))
        self.assertFalse(self.tool.current("review_visual"))

    def test_incorrect_first_half_pdf_is_rejected(self):
        self.rendered_fixture(sample_pages=3)
        with self.assertRaisesRegex(ValueError, "first-half sample"):
            self.tool.finalize()

    def test_document_publication_uses_document_upload_fields(self):
        for relative in self.tool.outputs:
            write_text(self.tool.workspace / relative, "Upload transaction fixture only")
        order = json.loads(self.snapshot.read_text(encoding="utf-8"))
        with patch.object(self.tool, "finalize"), \
             patch("codex_order_tools.api.fetch_order_read_only", return_value=order), \
             patch("codex_order_tools.api.claim_oldest"), patch("codex_order_tools.api.heartbeat"), \
             patch("codex_order_tools.api.sync_review_scores", return_value={"reviews": []}), \
             patch("codex_order_tools.api.submit_final") as upload:
            self.assertTrue(self.tool.publish()["published"])
            self.assertEqual(set(upload.call_args.args[2]), {"docx_file", "pdf_file", "sample_file", "sample_pdf_file"})
            self.assertTrue(self.tool.publish()["reused_receipt"])
            self.assertEqual(upload.call_count, 1)


if __name__ == "__main__":
    unittest.main()
