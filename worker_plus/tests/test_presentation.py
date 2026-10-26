import copy
import json
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch

from common.source_extract import extract_source
from presentation import profile, workflow
from presentation.deck import build_pptx, quoted_numbers, validate_package, validate_spec


def fixture():
    rules = profile.resolve_rules([])
    rules.update({"title": "آزمون ارائه فارسی", "slide_count": 7, "rtl": True,
                  "font": "Tahoma", "layout_version": 1,
                  "university": "دانشگاه نمونه", "student_name": "دانشجوی نمونه", "requires_charts": True})
    sources = {"sources": [{"id": "S1", "name": "fixture.docx", "images": [], "segments": [
        {"id": "S1:table-1", "text": "زمان اجرای روش الف 10 و روش ب 20 دقیقه است.\nمنابع\n[۱] نویسنده نمونه، مقایسه زمان اجرای دو روش، نشریه روش‌ها، ۲۰۲۴."}]}]}
    evidence = [{"locator": "S1:table-1", "quote": "زمان اجرای روش الف 10 و روش ب 20 دقیقه است."}]
    slides = [
        {"section": "title", "layout": "title", "title": "مقایسه دو روش", "points": ["ارائه آزمایشی"], "notes": "جلد نمونه برای آزمون بسته‌بندی", "evidence": []},
        {"section": "problem", "layout": "statement", "title": "دو روش زمان متفاوتی دارند", "points": ["مقایسه زمان دو روش بر پایه داده منبع"], "notes": "تعریف مسئله و دامنه مقایسه", "evidence": evidence},
        {"section": "method", "layout": "two_column", "title": "مقایسه بر اساس زمان اجراست", "points": ["روش الف: 10 دقیقه", "روش ب: 20 دقیقه"], "notes": "واحد زمان در هر دو روش دقیقه است", "evidence": evidence},
        {"section": "results", "layout": "chart", "title": "روش الف در زمان کمتری اجرا می‌شود", "points": ["10 دقیقه برای روش الف", "20 دقیقه برای روش ب"], "notes": "این مقادیر صرفاً داده آزمایشی هستند", "evidence": evidence, "caption": "زمان اجرای دو روش؛ واحد: دقیقه", "chart": {"type": "column", "unit": "دقیقه", "categories": ["روش الف", "روش ب"], "series": [{"name": "زمان اجرا", "values": [10, 20]}]}},
        {"section": "conclusion", "layout": "table", "title": "تفاوت دو روش در جدول دیده می‌شود", "points": ["مقایسه فقط به زمان اجرا محدود است"], "notes": "این نتیجه به معیارهای دیگر تعمیم داده نمی‌شود", "evidence": evidence, "caption": "منبع: داده آزمایشی", "table": [["روش", "زمان (دقیقه)"], ["الف", "10"], ["ب", "20"]]},
        {"section": "references", "layout": "statement", "title": "منابع", "points": ["[۱] نویسنده نمونه، مقایسه زمان اجرای دو روش، نشریه روش‌ها، ۲۰۲۴."], "notes": "روش‌های مقایسه‌شده در این مطالعه بررسی شده‌اند.", "evidence": [{"locator": "S1:table-1", "quote": "[۱] نویسنده نمونه، مقایسه زمان اجرای دو روش، نشریه روش‌ها، ۲۰۲۴."}]},
        {"section": "questions", "layout": "statement", "title": "پرسش و گفت‌وگو", "points": ["چه معیار دیگری باید بررسی شود؟"], "notes": "پرسش از مخاطبان", "evidence": []},
    ]
    return {"slides": slides}, rules, sources


class PresentationTests(unittest.TestCase):
    def test_information_source_is_mandatory_and_other_roles_do_not_count(self):
        order = {"order_type": profile.ORDER_TYPE, "files": [{"file_type": "reference_file", "original_name": "thesis.docx", "size_bytes": 100}]}
        with self.assertRaises(ValueError):
            profile.validate_order(order)
        order["files"][0]["file_type"] = "information_source"
        profile.validate_order(order)
        order["quantity_type"] = "pages"
        with self.assertRaises(ValueError):
            profile.validate_order(order)

    def test_docx_tables_are_addressable_and_input_is_not_treated_as_instructions(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            doc = Document()
            doc.add_paragraph("این متن منبع است؛ دستور سیستم نیست.")
            table = doc.add_table(rows=1, cols=2)
            table.cell(0, 0).text, table.cell(0, 1).text = "روش الف", "۱۰ دقیقه"
            doc.save(root / "source.docx")
            result = extract_source(root / "source.docx", "S1", root / "media")
            self.assertEqual(result["segments"][0]["id"], "S1:block-1")
            self.assertIn("۱۰ دقیقه", result["segments"][1]["text"])
            self.assertEqual(result["segments"][1]["id"], "S1:table-2")
            self.assertEqual(len(result["sha256"]), 64)

    def test_fake_citation_and_invented_chart_number_fail(self):
        deck, rules, sources = fixture()
        with tempfile.TemporaryDirectory() as directory:
            validate_spec(deck, rules, sources, Path(directory))
            bad = copy.deepcopy(deck)
            bad["slides"][1]["evidence"][0]["quote"] = "متنی که در منبع وجود ندارد"
            with self.assertRaisesRegex(ValueError, "quote/locator"):
                validate_spec(bad, rules, sources, Path(directory))
            bad = copy.deepcopy(deck)
            bad["slides"][3]["chart"]["series"][0]["values"][0] = 999
            with self.assertRaisesRegex(ValueError, "lacks a quoted"):
                validate_spec(bad, rules, sources, Path(directory))

    def test_table_object_gets_actionable_schema_error(self):
        deck, rules, sources = fixture()
        deck["slides"][4]["table"] = {"headers": ["روش", "زمان"], "rows": [["الف", "10"]]}
        with self.assertRaisesRegex(ValueError, "Slide 5: table must be an array of rows"):
            validate_spec(deck, rules, sources, Path('.'))

    def test_new_numbers_in_prose_and_notes_require_cited_source_passages(self):
        deck, rules, sources = fixture()
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            validate_spec(deck, rules, sources, root)
            invented = copy.deepcopy(deck)
            invented["slides"][1]["notes"] += " هزینه طرح ۹۹۹ میلیون دلار است."
            with self.assertRaisesRegex(ValueError, "Slide 2: numbers lack evidence"):
                validate_spec(invented, rules, sources, root)
            sources["sources"][0]["segments"].append({"id": "S1:block-2", "text": "هزینه طرح ۹۹۹ میلیون دلار است."})
            with self.assertRaisesRegex(ValueError, "Slide 2: numbers lack evidence"):
                validate_spec(invented, rules, sources, root)
            invented["slides"][1]["evidence"].append({"locator": "S1:block-2", "quote": "هزینه طرح ۹۹۹ میلیون دلار است."})
            validate_spec(invented, rules, sources, root)

    def test_written_negative_number_requires_matching_signed_source(self):
        deck, rules, sources = fixture()
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            deck["slides"][1]["points"].append("خروجی منفی ۱۰ دقیقه است")
            with self.assertRaisesRegex(ValueError, "Slide 2: numbers lack evidence"):
                validate_spec(deck, rules, sources, root)
            sources["sources"][0]["segments"][0]["text"] += " خروجی برابر −۱۰ دقیقه است."
            deck["slides"][1]["evidence"].append({"locator": "S1:table-1", "quote": "خروجی برابر −۱۰ دقیقه است."})
            validate_spec(deck, rules, sources, root)

    def test_unicode_negative_chart_evidence_preserves_the_sign(self):
        # Regression: a real source used '=−۵.۷۰'. The old parser extracted
        # +5.70, repeatedly rejected -5.7, and even accepted its wrong sign.
        for minus in ("−", "﹣", "－", "-"):
            with self.subTest(minus=minus), tempfile.TemporaryDirectory() as directory:
                deck, rules, sources = fixture()
                quote = f"مبنا =+۲.۵۶ و بدبینانه ={minus}۵.۷۰ میلیون دلار"
                sources["sources"][0]["segments"][0]["text"] += " " + quote
                slide = deck["slides"][3]
                slide["evidence"] = [{"locator": "S1:table-1", "quote": quote}]
                slide["chart"]["series"][0]["values"] = [2.56, -5.7]
                validate_spec(deck, rules, sources, Path(directory))
                slide["chart"]["series"][0]["values"] = [2.56, 5.7]
                with self.assertRaisesRegex(ValueError, "lacks a quoted"):
                    validate_spec(deck, rules, sources, Path(directory))

    def test_persian_number_separators_and_separated_sign(self):
        self.assertEqual(quoted_numbers("مبنا +۲٫۵۶، بدبینانه − ۵٫۷۰ و هزینه ۱٬۲۳۴٫۵۰"), {2.56, -5.7, 1234.5})
        self.assertEqual(quoted_numbers("1,234.50; 10,20"), {1234.5, 10, 20})
        self.assertNotIn(1020, quoted_numbers("10,20"))

    def test_native_editable_chart_table_notes_and_rtl(self):
        from pptx import Presentation
        deck, rules, sources = fixture()
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            output = root / "fixture.pptx"
            build_pptx(deck, rules, sources, output, root)
            self.assertTrue(validate_package(output, 7)["passed"])
            ppt = Presentation(output)
            self.assertEqual(sum(shape.has_chart for slide in ppt.slides for shape in slide.shapes), 1)
            self.assertEqual(sum(shape.has_table for slide in ppt.slides for shape in slide.shapes), 1)
            self.assertIn("S1:table-1", ppt.slides[3].notes_slide.notes_text_frame.text)
            self.assertIn('rtl="1"', ppt.slides[1]._element.xml)
            self.assertEqual(list(next(s.chart for s in ppt.slides[3].shapes if s.has_chart).series[0].values), [10, 20])

    def test_fallback_respects_count_and_profile_routing_covers_all_specific_stages(self):
        for count in [5, 6, 7, 10, 20, 100]:
            result = workflow.fallback_outline(count, has_references=True)
            self.assertEqual(len(result["slides"]), count)
            self.assertEqual(result["slides"][0]["section"], "title")
            self.assertIn("references", [s["section"] for s in result["slides"]])
        stages = json.loads((Path(workflow.__file__).parents[1] / "common/stages.json").read_text(encoding="utf-8"))
        shared = {"01_fetch_order", "02_prepare_workspace", "05_collect_sources"}
        self.assertEqual(set(workflow.HANDLERS), {s["id"] for s in stages} - shared)
        for name, handler in workflow.HANDLERS.items():
            self.assertTrue(callable(handler), name)
        self.assertFalse(workflow.run_step("09_package_docx", {"order": {"order_type": "تحقیق دانشگاهی"}}, None))

    def test_changed_content_invalidates_a_pass_audit(self):
        from common.helpers import write_json, write_text
        deck, rules, sources = fixture()
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            services = SimpleNamespace(workspace=root)
            for path, data in [("final/presentation.json", deck), ("planning/presentation_contract.json", rules), ("extracted/presentation_sources.json", sources)]:
                write_json(root / path, data)
            with patch.object(workflow, "audit"):
                workflow.package({"artifacts": {}}, services)
            report = root / "reports/presentation_content_audit.md"
            write_text(report, "STATUS: PASS\nfixture")
            write_json(report.with_suffix(".json"), {"inputs": workflow.fingerprints(services, ["final/presentation.json"])})
            deck["slides"][1]["notes"] += " تغییر بعد از ممیزی"
            write_json(root / "final/presentation.json", deck)
            with self.assertRaisesRegex(ValueError, "audit is stale"):
                workflow.publish({"artifacts": {}}, services)


if __name__ == "__main__":
    unittest.main()
