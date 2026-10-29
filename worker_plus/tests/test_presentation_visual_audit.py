import json
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import unittest

from common.helpers import write_json, write_text
from presentation import workflow
from presentation.recovery import PresentationVisualAuditError
from presentation.visual_audit import aggregate_visual_audit
from presentation.visual_checks import prepare_visual_evidence


class PresentationVisualAuditTests(unittest.TestCase):
    def make_fixture(self, root, damage=None):
        import fitz
        deck = {"slides": [{"layout": "statement", "evidence": [{"locator": "S1:block-12", "quote": "Slide body"}]}]}
        rules = {"font": "Helvetica", "slide_count": 1}
        write_json(root / "final/presentation.json", deck)
        write_json(root / "planning/presentation_contract.json", rules)
        sources = {"sources": [{"id": "S1", "name": "study.docx", "pagination": {}, "images": [],
                                "segments": [{"id": "S1:block-12", "text": "Slide body", "pages": [2]}]}]}
        write_json(root / "extracted/presentation_sources.json", sources)
        from pptx import Presentation
        from presentation.source_comments import add_source_comments, expected_comments
        prs = Presentation(); prs.slides.add_slide(prs.slide_layouts[6]); prs.save(root / "final/final.pptx")
        add_source_comments(root / "final/final.pptx", expected_comments(deck, sources, rules))
        with fitz.open() as pdf:
            page = pdf.new_page(width=960, height=540)
            page.insert_text((60, 60), "Slide body", fontsize=22)
            if damage:
                page.insert_text((60, 510), "S1:block-12", fontsize=13, render_mode=3 if damage == "invisible" else 0)
            if damage in {"white_cover", "dark_cover", "partial_cover"}:
                width = 3 if damage == "partial_cover" else 40
                color = (0, 0, 0) if damage == "dark_cover" else (1, 1, 1)
                page.draw_rect(fitz.Rect(60, 493, 60 + width, 513), fill=color, color=color, overlay=True)
            pdf.save(root / "final/final.pdf")
        return deck, rules

    def claim(self, **changes):
        return {"slide": 1, "region": "footer", "kind": "citation_clipped", "locator": "S1:block-12", "message": "Only the final 12 is visible", **changes}

    def model(self, issues, status="FAIL", reviewed=None):
        return f"STATUS: {status}\nISSUES_JSON: " + json.dumps({"issues": issues, "reviewed_slides": [1] if reviewed is None else reviewed})

    def test_searchable_text_with_missing_pixels_is_rejected(self):
        import fitz
        for damage in ("invisible", "white_cover", "dark_cover", "partial_cover"):
            with self.subTest(damage=damage), TemporaryDirectory() as directory:
                root = Path(directory)
                deck, rules = self.make_fixture(root, damage)
                with fitz.open(root / "final/final.pdf") as pdf:
                    self.assertIn("S1:block-12", pdf[0].get_text(), "Text extraction alone would incorrectly approve this")
                with self.assertRaises(PresentationVisualAuditError):
                    prepare_visual_evidence(root, deck, rules)
                report = json.loads((root / "reports/visual_evidence/checks.json").read_text())
                self.assertFalse(report["passed"])
                self.assertTrue(report["slides"][0]["visible_internal_references"])

    def test_comments_are_verified_and_missing_visible_footer_is_not_required(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            deck, rules = self.make_fixture(root)
            proof = prepare_visual_evidence(root, deck, rules)
            self.assertTrue(proof["passed"])
            self.assertEqual([], proof["slides"][0]["raster"]["locators"])
            self.assertTrue(proof["provenance"]["passed"])
            result = aggregate_visual_audit(root, self.model([], status="PASS"), deck, rules)
            self.assertTrue(result["passed"])
            self.assertEqual(0, len(result["resolved_claims"]))
            self.assertEqual("STATUS: PASS", result["model_status"])

    def test_body_chart_unknown_or_subjective_issues_still_fail(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            deck, rules = self.make_fixture(root)
            prepare_visual_evidence(root, deck, rules)
            for issue in (self.claim(region="body"), self.claim(kind="chart_color"), self.claim(kind="citation_unreadable"),
                          self.claim(locator="S1:block-99"), {"slide": 1, "message": "legacy unclassified defect"}):
                with self.subTest(issue=issue):
                    result = aggregate_visual_audit(root, self.model([issue], status="PASS"), deck, rules)
                    self.assertFalse(result["passed"], "A model PASS cannot hide an unresolved defect")
                    self.assertEqual([issue], result["issues"])

    def test_stale_pdf_spec_or_rules_cannot_resolve_model_claims(self):
        for changed in ("pdf", "pptx", "spec", "rules"):
            with self.subTest(changed=changed), TemporaryDirectory() as directory:
                root = Path(directory)
                deck, rules = self.make_fixture(root)
                prepare_visual_evidence(root, deck, rules)
                if changed == "pdf":
                    with (root / "final/final.pdf").open("ab") as stream:
                        stream.write(b"\n%changed\n")
                elif changed == "pptx":
                    with (root / "final/final.pptx").open("ab") as stream:
                        stream.write(b"changed comments")
                elif changed == "spec":
                    deck["slides"][0]["evidence"][0]["locator"] = "S1:block-99"
                else:
                    rules["font"] = "Tahoma"
                with self.assertRaises(ValueError):
                    aggregate_visual_audit(root, self.model([self.claim()]), deck, rules)

    def test_unreviewed_slides_and_ambiguous_fail_are_not_approved(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            deck, rules = self.make_fixture(root)
            prepare_visual_evidence(root, deck, rules)
            for text in (self.model([], reviewed=[]), self.model([], reviewed=[1, 1]), self.model([]),
                         'STATUS: PASS\nISSUES_JSON: {"issues":[]}'):
                with self.subTest(text=text), self.assertRaises(ValueError):
                    aggregate_visual_audit(root, text, deck, rules)

    def test_workflow_preserves_raw_fail_and_stamps_the_actual_decision_evidence(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            deck, rules = self.make_fixture(root)
            prepare_visual_evidence(root, deck, rules)
            raw = self.model([], status="PASS")
            def respond(prompt, target):
                self.assertEqual("presentation_visual_model_audit.md", target.name)
                self.assertIn("reviewed_slides", prompt)
                write_text(target, raw)
            workflow.audit(SimpleNamespace(workspace=root, run_codex=respond), "presentation_visual_audit.md", "fixture")
            self.assertEqual(raw, (root / "reports/presentation_visual_model_audit.md").read_text())
            self.assertTrue((root / "reports/presentation_visual_audit.md").read_text().startswith("STATUS: PASS\n"))
            decision = json.loads((root / "reports/presentation_visual_decision.json").read_text())
            self.assertEqual("STATUS: PASS", decision["model_status"])
            self.assertEqual([], decision["resolved_claims"])
            stamp = json.loads((root / "reports/presentation_visual_audit.json").read_text())
            self.assertIn("reports/visual_evidence/checks.json", stamp["inputs"])
            self.assertIn("reports/presentation_visual_model_audit.md", stamp["inputs"])
            self.assertIn("reports/presentation_visual_decision.json", stamp["inputs"])


if __name__ == "__main__":
    unittest.main()
