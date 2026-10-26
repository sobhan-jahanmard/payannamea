from copy import deepcopy
import json
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import unittest
from unittest.mock import Mock, patch

from common.helpers import write_json, write_text
from common import runner
from common.order_context import load, new_context, save
from presentation import workflow
from presentation.deck import build_pptx, quoted_numbers, validate_spec
from presentation.recovery import PresentationAuditError, PresentationIntegrityError, PresentationNoProgressError, apply_slide_repairs, audit_issues, verified_content_checkpoint
from tests.test_presentation import fixture


class PresentationRecoveryTests(unittest.TestCase):
    def prepare(self, root):
        deck, rules, sources = fixture()
        rules["cover_lines"] = deck["slides"][0]["points"]
        for path, data in [("final/presentation.json", deck), ("planning/presentation_contract.json", rules), ("extracted/presentation_sources.json", sources)]:
            write_json(root / path, data)
        return deck, rules, sources

    def test_summary_or_zero_slides_never_overwrites_the_previous_draft(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            deck, rules, sources = self.prepare(root)
            original = (root / "final/presentation.json").read_bytes()
            for reply in ({}, {"summary": "fixed"}, {"slides": []}):
                def respond(prompt, target):
                    self.assertNotEqual(target, root / "final/presentation.json")
                    write_json(target, reply)
                service = SimpleNamespace(workspace=root, run_codex=respond)
                def validate(value):
                    validate_spec(value, rules, sources, root)
                    return value
                with self.assertRaises((ValueError, KeyError)):
                    workflow.model_json(service, "fixture", "final/presentation.json", transform=validate)
                self.assertEqual(original, (root / "final/presentation.json").read_bytes())

    def test_model_source_edit_is_rolled_back_and_not_accepted(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            deck, _, _ = self.prepare(root)
            source = root / "extracted/presentation_sources.json"
            original = source.read_bytes()
            def respond(prompt, target):
                source.write_text('{"sources":[]}', encoding="utf-8")
                write_json(target, deck)
            service = SimpleNamespace(workspace=root, run_codex=respond)
            with self.assertRaises(PresentationIntegrityError):
                workflow.model_json(service, "fixture", "final/presentation.json")
            self.assertEqual(original, source.read_bytes())

    def test_repair_only_changes_allowed_slides_and_rejects_no_progress(self):
        deck, _, _ = fixture()
        changed = deepcopy(deck["slides"][1])
        changed["notes"] += " corrected"
        result = apply_slide_repairs(deck, {"patches": [{"slide_number": 2, "slide": changed}]}, [2])
        self.assertEqual(deck["slides"][:1], result["slides"][:1])
        self.assertEqual(deck["slides"][2:], result["slides"][2:])
        with self.assertRaises(ValueError):
            apply_slide_repairs(deck, {"patches": [{"slide_number": 3, "slide": changed}]}, [2])
        with self.assertRaises(PresentationNoProgressError):
            apply_slide_repairs(deck, {"patches": [{"slide_number": 2, "slide": deck["slides"][1]}]}, [2])

    def test_rerunning_step_13_reuses_valid_saved_draft_and_still_requires_audits(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            deck, _, _ = self.prepare(root)
            write_json(root / "drafts/presentation_last_valid.json", deck)
            feedback = {"failed_step": "07_generate_content", "error": "Presentation repair made no progress; current draft preserved"}
            service = SimpleNamespace(workspace=root, recovery_feedback=feedback, run_codex=Mock(side_effect=AssertionError("Must not regenerate valid slides")))
            context = {"artifacts": {}, "recovery_history": [feedback]}
            workflow.generate(context, service)
            workflow.generate(context, service)
            service.run_codex.assert_not_called()
            self.assertEqual(deck, json.loads((root / "final/presentation.json").read_text(encoding="utf-8")))
            self.assertTrue(json.loads((root / "reports/stage_checks/presentation_generation.json").read_text(encoding="utf-8"))["content_and_source_audits_required"])
            self.assertFalse((root / "reports/presentation_source_audit.json").exists())

    def test_usage_pause_keeps_current_failed_audit_for_targeted_repair(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            deck, _, _ = self.prepare(root)
            inputs = workflow.fingerprints(SimpleNamespace(workspace=root),
                ["final/presentation.json", "planning/presentation_contract.json", "extracted/presentation_sources.json"])
            write_json(root / "reports/presentation_content_audit.json", {"passed": False, "inputs": inputs})
            write_text(root / "reports/presentation_content_audit.md",
                'STATUS: FAIL\nISSUES_JSON: {"issues":[{"slide":2,"message":"source detail missing"}]}')
            audit_failure = {"failed_step": "08_review_content", "report_path": "reports/presentation_content_audit.md",
                             "repair_slides": [2], "error": "source detail missing"}
            quota_failure = {"failed_step": "07_generate_content", "error": "usage limit"}
            changed = deepcopy(deck["slides"][1])
            changed["notes"] += " More source detail."
            def respond(prompt, target):
                self.assertIn('"patches"', prompt)
                self.assertIn('source detail missing', prompt)
                write_json(target, {"patches": [{"slide_number": 2, "slide": changed}]})
            service = SimpleNamespace(workspace=root, recovery_feedback=quota_failure,
                                      run_codex=Mock(side_effect=respond))
            workflow.generate({"artifacts": {}, "recovery_history": [audit_failure, quota_failure]}, service)
            saved = json.loads((root / "final/presentation.json").read_text(encoding="utf-8"))
            self.assertEqual(saved["slides"][1], changed)
            self.assertEqual(saved["slides"][2:], deck["slides"][2:])

    def test_identical_patch_gets_one_fresh_audit_then_stops_if_still_rejected(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            deck, _, _ = self.prepare(root)
            feedback = {"failed_step": "08b_audit_source_compliance", "error": "Slide 2: unsupported claim", "repair_slides": [2]}
            inputs = workflow.fingerprints(SimpleNamespace(workspace=root), ["final/presentation.json", "planning/presentation_contract.json", "extracted/presentation_sources.json"])
            write_json(root / "reports/presentation_source_audit.json", {"passed": False, "inputs": inputs})
            def respond(prompt, target):
                write_json(target, {"patches": [{"slide_number": 2, "slide": deck["slides"][1]}]})
            service = SimpleNamespace(workspace=root, recovery_feedback=feedback, run_codex=Mock(side_effect=respond))
            context = {"artifacts": {}, "recovery_history": [feedback]}
            workflow.generate(context, service)
            self.assertFalse(json.loads((root / "reports/presentation_source_audit.json").read_text(encoding="utf-8"))["passed"])
            with self.assertRaisesRegex(PresentationNoProgressError, "fresh audit"):
                workflow.generate(context, service)
            self.assertEqual(2, service.run_codex.call_count)

    def test_changed_inputs_do_not_reuse_an_old_generation_checkpoint(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            deck, _, sources = self.prepare(root)
            context = {"artifacts": {}}
            workflow.finish_generation(context, SimpleNamespace(workspace=root), deck)
            sources["sources"][0]["segments"][0]["text"] += " Additional source information."
            write_json(root / "extracted/presentation_sources.json", sources)
            def respond(prompt, target):
                write_json(target, deck)
            service = SimpleNamespace(workspace=root, recovery_feedback={"failed_step": "07_generate_content", "error": "No progress"}, run_codex=Mock(side_effect=respond))
            workflow.generate(context, service)
            service.run_codex.assert_called_once()

    def test_legacy_report_only_targets_the_unresolved_slide(self):
        text = "STATUS: FAIL\n- اسلاید ۱۵: خطای قبلی واقعاً رفع شده است؛ مقدار تطابق کامل دارد.\n- اسلاید ۱۸: ادعا توسط شاهد پشتیبانی نمی‌شود؛ ردیابی ناقص است."
        self.assertEqual([18], [issue["slide"] for issue in audit_issues(text, 20)])

    def test_rtl_postfix_minus_does_not_turn_a_loss_positive(self):
        self.assertEqual({2.56, -5.7}, quoted_numbers("مبنا ۲٫۵۶ و بدبینانه ۵٫۷۰− میلیون دلار"))
        self.assertEqual({6.9, -5.7}, quoted_numbers("۶.۹-۵.۷"))

    def test_audit_error_exposes_report_and_exact_issue_for_repair(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            self.prepare(root)
            def respond(prompt, target):
                self.assertIn("قبل از بسته‌بندی", prompt)
                write_text(target, 'STATUS: FAIL\nISSUES_JSON: {"issues":[{"slide":2,"message":"Use the evidence at S1:table-1"}]}')
            service = SimpleNamespace(workspace=root, run_codex=respond)
            with self.assertRaises(PresentationAuditError) as failure:
                workflow.audit(service, "fixture_audit.md", "check fixture")
            self.assertEqual([2], failure.exception.repair_slides)
            self.assertEqual("reports/fixture_audit.md", failure.exception.report_path)
            self.assertIn("S1:table-1", str(failure.exception))

    def test_visual_auditor_checks_packaged_outputs_and_cannot_edit_evidence(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            self.prepare(root)
            pdf = root / "final/final.pdf"
            pdf.write_bytes(b"original-rendered-pdf")
            image = root / "reports/visual_evidence/slide-001-footer.png"
            image.parent.mkdir(parents=True)
            image.write_bytes(b"original-footer-image")
            def respond(prompt, target):
                self.assertIn("پس از ساخت", prompt)
                self.assertNotIn("قبل از بسته‌بندی", prompt)
                pdf.write_bytes(b"modified")
                image.unlink()
                write_text(target, 'STATUS: PASS\nISSUES_JSON: {"issues":[]}')
            service = SimpleNamespace(workspace=root, run_codex=respond)
            with self.assertRaises(PresentationIntegrityError):
                workflow.audit(service, "presentation_visual_audit.md", "check fixture")
            self.assertEqual(b"original-rendered-pdf", pdf.read_bytes())
            self.assertEqual(b"original-footer-image", image.read_bytes())
            self.assertFalse((root / "reports/presentation_visual_audit.json").exists())

    def test_runner_retries_audit_with_one_slide_patch_and_keeps_all_other_slides(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            workspace = root / "in_progress"
            original, _, _ = self.prepare(workspace)
            context = new_context(workspace, "full")
            context.update(order_id="isolated-presentation-repair", order={"order_type": "ارائه و پاورپوینت"},
                           completed_steps=runner.STEPS[:runner.STEPS.index("08b_audit_source_compliance")], status="in_progress")
            save(workspace, context)
            service = Mock(workspace=workspace)
            service.recovery_feedback = None
            service.cleanup_run.return_value = {"released_docx_targets": [], "terminated_automation_word_pids": []}
            calls = {"source": 0, "repair": 0}
            def respond(prompt, target):
                if target.name == "presentation_source_audit.md":
                    calls["source"] += 1
                    write_text(target, 'STATUS: FAIL\nISSUES_JSON: {"issues":[{"slide":2,"message":"Add the missing source locator S1:table-1"}]}' if calls["source"] == 1 else 'STATUS: PASS\nISSUES_JSON: {"issues":[]}')
                elif target.name == "presentation_content_audit.md":
                    write_text(target, 'STATUS: PASS\nISSUES_JSON: {"issues":[]}')
                else:
                    calls["repair"] += 1
                    self.assertIn("Add the missing source locator", prompt)
                    corrected = deepcopy(original["slides"][1])
                    corrected["caption"] = "Source: S1:table-1"
                    write_json(target, {"patches": [{"slide_number": 2, "slide": corrected}]})
            service.run_codex.side_effect = respond
            def complete(ctx, _services):
                ctx["status"] = "dry_run_complete"
            # Later package stages are deliberately stubbed: no order or Office run.
            handlers = {name: complete for name in runner.STEPS[runner.STEPS.index("09_package_docx"): ]}
            stage = Mock(TITLE="fixture")
            with patch.object(runner, "load_config", return_value=SimpleNamespace(workspace_root=root)), \
                 patch.object(runner, "Services", return_value=service), \
                 patch.object(runner.sys, "argv", ["run.py", "--resume", "--offline"]), \
                 patch.object(runner.importlib, "import_module", return_value=stage), \
                 patch.dict(workflow.HANDLERS, handlers):
                runner.main()
            saved = load(workspace)
            current = json.loads((workspace / "final/presentation.json").read_text(encoding="utf-8"))
            self.assertEqual({"source": 2, "repair": 1}, calls)
            self.assertEqual(original["slides"][:1], current["slides"][:1])
            self.assertEqual(original["slides"][2:], current["slides"][2:])
            self.assertEqual([2], saved["recovery_history"][-1]["repair_slides"])
            self.assertEqual("dry_run_complete", saved["status"])
            self.assertEqual(1, len([node for node in saved["execution_trace"]["nodes"] if node["status"] == "failed"]))

    def test_negative_chart_fill_is_not_inverted_to_white(self):
        from pptx import Presentation
        with TemporaryDirectory() as directory:
            root = Path(directory)
            deck, rules, sources = self.prepare(root)
            deck["slides"][3]["chart"]["series"][0]["values"] = [2.56, -5.7]
            build_pptx(deck, rules, sources, root / "fixture.pptx", root)
            slide = Presentation(root / "fixture.pptx").slides[3]
            chart = next(shape.chart for shape in slide.shapes if shape.has_chart)
            self.assertFalse(chart.series[0].invert_if_negative)
            self.assertEqual([2.56, -5.7], list(chart.series[0].values))

    def test_mixed_language_runs_and_reference_names_keep_their_direction(self):
        from pptx import Presentation
        with TemporaryDirectory() as directory:
            root = Path(directory)
            deck, rules, sources = self.prepare(root)
            deck["slides"][1]["points"] = ["خرید ناخالص PET: ۳۶۹۱ تن/سال"]
            deck["slides"][5]["points"] = ["Amaplast؛ Process Engineering Associates؛ JRC؛ مقررات ایران", "WITS/UN Comtrade، ۲۰۲۳"]
            before = deepcopy(deck)
            build_pptx(deck, rules, sources, root / "fixture.pptx", root)
            ppt = Presentation(root / "fixture.pptx")
            mixed = next(shape for shape in ppt.slides[1].shapes if shape.has_text_frame and "PET" in shape.text)
            self.assertIn("\u202aPET\u202c\u200f", mixed.text)
            self.assertIn('lang="en-US"', mixed._element.xml)
            paragraphs = [p for shape in ppt.slides[5].shapes if shape.has_text_frame for p in shape.text_frame.paragraphs]
            organization = next(p for p in paragraphs if p.text == "Process Engineering Associates")
            self.assertEqual("0", organization._p.get_or_add_pPr().get("rtl"))
            self.assertTrue(any(p.text == "WITS/UN Comtrade, ۲۰۲۳" for p in paragraphs))
            self.assertEqual(before, deck)

    def test_verified_content_checkpoint_rejects_stale_source_or_content(self):
        with TemporaryDirectory() as directory:
            root = Path(directory)
            self.prepare(root)
            service = SimpleNamespace(workspace=root)
            inputs = workflow.fingerprints(service, ["final/presentation.json", "extracted/presentation_sources.json", "planning/presentation_contract.json"])
            for name in ("presentation_content_audit", "presentation_source_audit"):
                write_text(root / f"reports/{name}.md", "STATUS: PASS")
                write_json(root / f"reports/{name}.json", {"passed": True, "inputs": inputs})
            self.assertTrue(verified_content_checkpoint(service))
            write_json(root / "extracted/presentation_sources.json", {"sources": []})
            self.assertFalse(verified_content_checkpoint(service))

    def test_rendering_recovery_never_regenerates_audited_content(self):
        for legacy_step_13 in (False, True):
            with self.subTest(legacy_step_13=legacy_step_13), TemporaryDirectory() as directory:
                root = Path(directory)
                workspace = root / "in_progress"
                _, rules, _ = self.prepare(workspace)
                rules["font"] = "Helvetica"
                write_json(workspace / "planning/presentation_contract.json", rules)
                original = (workspace / "final/presentation.json").read_bytes()
                service = Mock(workspace=workspace)
                service.recovery_feedback = None
                service.cleanup_run.return_value = {"released_docx_targets": [], "terminated_automation_word_pids": []}
                inputs = workflow.fingerprints(service, ["final/presentation.json", "extracted/presentation_sources.json", "planning/presentation_contract.json"])
                for name in ("presentation_content_audit", "presentation_source_audit"):
                    write_text(workspace / f"reports/{name}.md", "STATUS: PASS")
                    write_json(workspace / f"reports/{name}.json", {"passed": True, "inputs": inputs})
                context = new_context(workspace, "full")
                context.update(order_id="isolated-rendering-repair", order={"order_type": "ارائه و پاورپوینت"}, status="in_progress",
                               completed_steps=runner.STEPS[:12 if legacy_step_13 else 20])
                if legacy_step_13:
                    context["recovery_history"] = [{"failed_step": "10c_verify_persian_pagination", "restart_step": "07_generate_content", "error": "visual rendering fault"}]
                save(workspace, context)
                audit_calls = 0
                def respond(prompt, target):
                    nonlocal audit_calls
                    self.assertEqual("presentation_visual_model_audit.md", target.name, "Generation must never run for a renderer defect")
                    audit_calls += 1
                    failed = not legacy_step_13 and audit_calls == 1
                    write_text(target, ('STATUS: FAIL' if failed else 'STATUS: PASS') + '\nISSUES_JSON: ' + json.dumps({"reviewed_slides": list(range(1, 8)), "issues": [{"slide": 4, "kind": "chart_color", "message": "negative bar is white"}] if failed else []}))
                service.run_codex.side_effect = respond
                def packaged(ctx, _services):
                    import fitz
                    (workspace / "final/final.pptx").write_bytes(b"isolated-package-fixture")
                    rules = workflow.contract(service)
                    with fitz.open() as pdf:
                        for slide in workflow.spec(service)["slides"]:
                            page = pdf.new_page(width=960, height=540)
                            page.insert_text((60, 60), "Fixture slide", fontname="Helvetica", fontsize=22)
                            footer = " . ".join(dict.fromkeys(e["locator"] for e in slide["evidence"]))
                            page.insert_text((60, 510), footer, fontname="Helvetica", fontsize=13)
                        pdf.save(workspace / "final/final.pdf")
                def visual(ctx, svc):
                    from presentation.visual_checks import prepare_visual_evidence
                    prepare_visual_evidence(workspace, workflow.spec(svc), workflow.contract(svc))
                    workflow.audit(svc, "presentation_visual_audit.md", "check fixture")
                def complete(ctx, _services):
                    ctx["status"] = "dry_run_complete"
                handlers = {name: packaged for name in runner.STEPS[16:20]}
                handlers.update({"10c_verify_persian_pagination": visual, "10_validate_and_publish": complete})
                packaged(context, service)
                with patch.object(runner, "load_config", return_value=SimpleNamespace(workspace_root=root)), \
                     patch.object(runner, "Services", return_value=service), \
                     patch.object(runner.sys, "argv", ["run.py", "--resume", "--offline"]), \
                     patch.object(runner.importlib, "import_module", return_value=Mock(TITLE="fixture")), \
                     patch.dict(workflow.HANDLERS, handlers):
                    if not legacy_step_13:
                        # An unchanged native renderer cannot fix itself. Keep
                        # the detailed failure and wait for a manual code fix.
                        with self.assertRaises(SystemExit):
                            runner.main()
                        stopped = load(workspace)
                        self.assertEqual(1, audit_calls)
                        self.assertEqual(runner.STEPS[:16], stopped["completed_steps"])
                        self.assertEqual(1, len(stopped["recovery_history"]))
                    runner.main()
                saved = load(workspace)
                self.assertEqual(original, (workspace / "final/presentation.json").read_bytes())
                self.assertEqual("09_package_docx", saved["recovery_history"][-1]["restart_step"])
                self.assertEqual("rendering", saved["recovery_history"][-1]["repair_kind"])
                self.assertEqual("dry_run_complete", saved["status"])
                self.assertEqual(runner.STEPS, saved["completed_steps"])

    def test_citations_show_every_locator_at_readable_size(self):
        from pptx import Presentation
        with TemporaryDirectory() as directory:
            root = Path(directory)
            deck, rules, sources = self.prepare(root)
            deck["slides"][1]["evidence"] = [{"locator": f"S1:block-{n}", "quote": "fixture"} for n in range(10, 16)]
            build_pptx(deck, rules, sources, root / "fixture.pptx", root)
            footer = next(shape for shape in Presentation(root / "fixture.pptx").slides[1].shapes if shape.name == "Source citations")
            self.assertEqual(" · ".join(f"S1:block-{n}" for n in range(10, 16)), footer.text)
            self.assertTrue(all(run.font.size.pt >= 13 for p in footer.text_frame.paragraphs for run in p.runs))

    def test_pdf_footer_checks_detect_missing_prefix_and_small_text(self):
        import fitz
        from presentation.visual_checks import prepare_visual_evidence
        from presentation.recovery import PresentationVisualAuditError
        for footer, size, passed in (("S1:block-12", 13, True), ("12", 13, False), ("S1:block-123", 13, False), ("S1:block-12", 11, False)):
            with self.subTest(footer=footer, size=size), TemporaryDirectory() as directory:
                root = Path(directory)
                (root / "final").mkdir()
                with fitz.open() as pdf:
                    page = pdf.new_page(width=960, height=540)
                    # The correct locator elsewhere cannot mask a broken footer.
                    page.insert_text((60, 60), "Body S1:block-12", fontsize=22)
                    page.insert_text((60, 510), footer, fontsize=size)
                    pdf.save(root / "final/final.pdf")
                deck = {"slides": [{"layout": "statement", "evidence": [{"locator": "S1:block-12"}]}]}
                rules = {"font": "Helvetica", "slide_count": 1}
                if passed:
                    result = prepare_visual_evidence(root, deck, rules)
                    self.assertTrue(result["passed"])
                else:
                    with self.assertRaises(PresentationVisualAuditError):
                        prepare_visual_evidence(root, deck, rules)
                self.assertTrue((root / "reports/visual_evidence/slide-001-footer.png").is_file())
                self.assertEqual(passed, json.loads((root / "reports/visual_evidence/checks.json").read_text())["passed"])


if __name__ == "__main__":
    unittest.main()
